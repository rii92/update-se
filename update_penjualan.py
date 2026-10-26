#!/usr/bin/env python3
"""
Batch Automation Update 5 Variabel Nilai & Komentar FASIH-SM SE2026
Mengubah 5 Variabel pada blok "KETERANGAN USAHA/PERUSAHAAN" (rincian Blok 27)
berdasarkan file Excel (source/penjualan.xlsx) hanya dengan 1 perintah:

1. Otomatis membuka Google Chrome dengan profil terisolasi (./chrome_data).
2. Otomatis login SSO BPS jika belum login / sesi habis.
3. Otomatis melewati (filter) baris yang sudah SUCCESS sebelumnya.
4. Untuk setiap baris:
   a. Membuka link kuesioner dalam mode edit (.../edit).
   b. Masuk ke section "KETERANGAN USAHA/PERUSAHAAN".
   c. Memilih sub-blok usaha sesuai "NAMA USAHA".
   d. Mengisi 5 variabel berikut (urut) + catatan pada masing-masing:
        1) Nilai penjualan barang dan jasa       (wrapper prefix: nilai_pendapatan)
        2) Biaya pembelian barang yang terjual   (wrapper prefix: biaya_pembelian)
        3) Biaya produksi                        (wrapper prefix: biaya_produksi)
        4) Biaya operasional                     (wrapper prefix: operasional)
        5) Upah/gaji                             (wrapper prefix: gaji)
        Komentar otomatis: "#1" jika nilai berubah, "#0" jika nilai sama/tidak berubah.
        Komentar manual per variabel bisa diisi lewat kolom remark_* di Excel.
   e. Menambah catatan pada field KBLI (prefix kbli#, fallback kbli_genai#) dengan
      default "#0" karena nilai KBLI tidak diubah.
   f. Submit assignment (masuk CATATAN -> switch anomali -> Kirim -> Konfirmasi).
5. Otomatis merekap hasil ke file Excel (source/penjualan_result.xlsx).

Kolom file sumber (source/penjualan.xlsx):
    link               : link assignment FASIH (boleh tanpa /edit)
    nama_usaha         : nama sub-blok usaha di section KETERANGAN USAHA/PERUSAHAAN
    nilai_penjualan    : 27.a Nilai penjualan barang dan jasa
    biaya_pembelian    : 27.b Biaya pembelian barang yang terjual
    biaya_produksi     : 27.c Biaya produksi
    biaya_operasional  : 27.d Biaya operasional
    upah               : 26.a Upah/gaji
    komentar           : (opsional/diabaikan) komentar kini otomatis #1/#0 per nilai

  Kolom opsional komentar manual per variabel (override otomatis #1/#0):
    remark_nilai_penjualan, remark_biaya_pembelian, remark_biaya_produksi,
    remark_biaya_operasional, remark_upah

  Kolom opsional TIMPA catatan lama (isi 1 = timpa, kosong/0 = pertahankan catatan lama):
    timpa_remark                       : berlaku untuk SEMUA variabel di baris itu
    timpa_remark_nilai_penjualan       : khusus variabel 27.a (prioritas lebih tinggi)
    timpa_remark_biaya_pembelian       : khusus 27.b
    timpa_remark_biaya_produksi        : khusus 27.c
    timpa_remark_biaya_operasional     : khusus 27.d
    timpa_remark_upah                  : khusus 26.a
    timpa_remark_kbli                  : khusus KBLI (default KBLI selalu ditimpa)

Contoh:
    python update_penjualan.py --init-sample      # buat contoh file Excel
    python update_penjualan.py                    # jalankan (skip SUCCESS)
    python update_penjualan.py --headless
    python update_penjualan.py --dry-run --limit 10
    python update_penjualan.py --start 1 --end 50
    python update_penjualan.py --force
"""

import argparse
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from playwright.sync_api import Page, sync_playwright
from rich.console import Console
from rich.panel import Panel
from rich.table import Table

import config

console = Console()

PROGRESS_FILE = Path("source/penjualan_progress.json")
RESULT_EXCEL_FILE = Path("source/penjualan_result.xlsx")
DEFAULT_SOURCE_FILE = Path("source/penjualan.xlsx")

# 5 Variabel yang diubah, urut.
# (kolom Excel, prefix id wrapper React, label, nama kolom remark manual)
VARIABEL_MAP = [
    ("nilai_penjualan", "nilai_pendapatan", "27.a Nilai penjualan barang dan jasa", "remark_nilai_penjualan"),
    ("biaya_pembelian", "biaya_pembelian", "27.b Biaya pembelian barang yang terjual", "remark_biaya_pembelian"),
    ("biaya_produksi", "biaya_produksi", "27.c Biaya produksi", "remark_biaya_produksi"),
    ("biaya_operasional", "operasional", "27.d Biaya operasional", "remark_biaya_operasional"),
    ("upah", "gaji", "26.a Upah/gaji", "remark_upah"),
]

REMARK_KEYS = [v[3] for v in VARIABEL_MAP]
# Key timpa per variabel + KBLI + global
TIMPA_KEYS = ["timpa_" + k for k in REMARK_KEYS] + ["timpa_remark_kbli", "timpa_remark"]

# Alias nama header Excel yang diterima untuk kolom remark manual
REMARK_ALIASES = {
    "remark_nilai_penjualan": ["remark_nilai_penjualan", "catatan_nilai_penjualan", "komentar_nilai_penjualan"],
    "remark_biaya_pembelian": ["remark_biaya_pembelian", "catatan_biaya_pembelian", "komentar_biaya_pembelian"],
    "remark_biaya_produksi": ["remark_biaya_produksi", "catatan_biaya_produksi", "komentar_biaya_produksi"],
    "remark_biaya_operasional": ["remark_biaya_operasional", "catatan_biaya_operasional", "komentar_biaya_operasional"],
    "remark_upah": ["remark_upah", "catatan_upah", "remark_gaji", "catatan_gaji", "komentar_upah"],
}

# JS bersama: gulir window + semua container scroll satu langkah
_SCROLL_STEP_JS = """(step) => {
    window.scrollBy({ top: step, behavior: "instant" });
    const scrollers = Array.from(document.querySelectorAll("div")).filter(el => {
        const oy = getComputedStyle(el).overflowY;
        return (oy === "auto" || oy === "scroll") && el.scrollHeight > el.clientHeight + 20;
    });
    scrollers.forEach(el => el.scrollBy({ top: step, behavior: "instant" }));
}"""

_HAS_MORE_SCROLL_JS = """(() => {
    const scrollers = Array.from(document.querySelectorAll("div")).filter(el => {
        const oy = getComputedStyle(el).overflowY;
        return (oy === "auto" || oy === "scroll") && el.scrollHeight > el.clientHeight + 20;
    });
    return scrollers.some(el => el.scrollTop + el.clientHeight < el.scrollHeight - 5);
})()"""

_CLOSE_DIALOG_JS = """(() => {
    const d = document.querySelector("[role='dialog']") || document.querySelector("[class*='max-w-lg']");
    const scope = d || document;
    const b = scope.querySelector("button[aria-label='Dismiss']") ||
              Array.from(scope.querySelectorAll("button")).find(x => {
                  const t = x.innerText.trim();
                  return t === "Close" || t === "Batal";
              });
    if (b) { b.click(); return true; }
    return false;
})()"""


# ─────────────────────────────────────────────────────────────────────────────
# Utilitas kecil
# ─────────────────────────────────────────────────────────────────────────────

def _is_true(val) -> bool:
    """True bila isi sel timpa bernilai 1 / true / ya / y (mendukung '1.0' dari Excel)."""
    if val is None:
        return False
    s = str(val).strip().lower()
    if s in ("1", "1.0", "true", "ya", "y", "yes"):
        return True
    return False


def _fmt_cell(val) -> str:
    """Ubah nilai sel Excel jadi string rapi. Float bulat (104160000.0) -> '104160000'."""
    if val is None:
        return ""
    if isinstance(val, bool):
        return str(val)
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def _normalize_num(val: Optional[str]) -> str:
    """Normalisasi angka: buang pemisah ribuan/desimal & spasi untuk perbandingan."""
    if val is None:
        return ""
    return str(val).strip().replace(".", "").replace(",", "").replace(" ", "")


# ─────────────────────────────────────────────────────────────────────────────
# Excel sumber / progress / hasil
# ─────────────────────────────────────────────────────────────────────────────

def create_sample_source(file_path: Path = DEFAULT_SOURCE_FILE):
    """Membuat file Excel sumber contoh jika belum ada."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Penjualan"
    ws.append([
        "link", "nama_usaha", "nilai_penjualan", "biaya_pembelian",
        "biaya_produksi", "biaya_operasional", "upah", "komentar",
        "remark_nilai_penjualan", "remark_biaya_pembelian", "remark_biaya_produksi",
        "remark_biaya_operasional", "remark_upah",
        "timpa_remark",
        "timpa_remark_nilai_penjualan", "timpa_remark_biaya_pembelian",
        "timpa_remark_biaya_produksi", "timpa_remark_biaya_operasional",
        "timpa_remark_upah", "timpa_remark_kbli",
    ])
    ws.append([
        config.DEFAULT_ASSIGNMENT_URL,
        "WARUNG SEMBAKO",
        "104160000",
        "83328000",
        "1812000",
        "4228000",
        "24000000",
        "#1",
        "#1", "#1", "#1", "#1", "#1",
        "0",
        "", "", "", "", "", "",
    ])
    file_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(file_path)
    console.print(f"[bold green]✔ Contoh file sumber dibuat:[/bold green] [yellow]{file_path}[/yellow]")
    console.print(
        "[dim]Kolom wajib: link, nama_usaha, nilai_penjualan, biaya_pembelian, "
        "biaya_produksi, biaya_operasional, upah. Satu baris per usaha.\n"
        "Opsional: remark_* (komentar manual), timpa_remark / timpa_remark_* (1 = timpa catatan lama).[/dim]"
    )


def load_progress() -> Dict[str, dict]:
    """Memuat data checkpoint dari file json dan file Excel hasil (jika ada)."""
    progress = {}

    if PROGRESS_FILE.exists():
        try:
            with open(PROGRESS_FILE, "r", encoding="utf-8") as f:
                progress = json.load(f)
        except Exception:
            pass

    if RESULT_EXCEL_FILE.exists():
        try:
            wb = openpyxl.load_workbook(RESULT_EXCEL_FILE, data_only=True)
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None) or []
            norm = [str(h).strip().lower() if h is not None else "" for h in header]

            def find_col(*keywords):
                for i, h in enumerate(norm):
                    if all(k in h for k in keywords):
                        return i
                return None

            c_no = find_col("no")
            c_link = find_col("link")
            c_nama = find_col("usaha")
            c_nilai = find_col("nilai")
            c_pembelian = find_col("pembelian")
            c_produksi = find_col("produksi")
            c_operasional = find_col("operasional")
            c_upah = find_col("upah")
            c_komentar = find_col("komentar")
            c_status = find_col("status")
            c_detail = find_col("detail")
            c_waktu = find_col("waktu")

            for row in rows:
                def get(c, default=""):
                    if c is None or len(row) <= c or row[c] is None:
                        return default
                    return str(row[c]).strip()

                idx_str = get(c_no)
                if not idx_str:
                    continue
                # Nomor dari Excel bisa terbaca "1.0"
                try:
                    idx_str = str(int(float(idx_str)))
                except ValueError:
                    pass
                status = get(c_status)
                if idx_str not in progress and status:
                    progress[idx_str] = {
                        "link": get(c_link),
                        "nama_usaha": get(c_nama),
                        "nilai_penjualan": get(c_nilai),
                        "biaya_pembelian": get(c_pembelian),
                        "biaya_produksi": get(c_produksi),
                        "biaya_operasional": get(c_operasional),
                        "upah": get(c_upah),
                        "komentar": get(c_komentar) or "#1",
                        "status": status,
                        "detail": get(c_detail),
                        "timestamp": get(c_waktu),
                    }
        except Exception:
            pass

    return progress


def save_progress(data: Dict[str, dict]):
    """Menyimpan data checkpoint ke file json secara atomik (aman bila proses terhenti)."""
    PROGRESS_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = PROGRESS_FILE.with_suffix(PROGRESS_FILE.suffix + ".tmp")
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp_path, PROGRESS_FILE)


def export_result_excel(tasks: List[dict], progress_data: Dict[str, dict], output_path: Path):
    """Mengekspor rekapitulasi status eksekusi ke Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hasil Update Penjualan"

    headers = [
        "No", "Link", "Nama Usaha", "Nilai Penjualan", "Biaya Pembelian",
        "Biaya Produksi", "Biaya Operasional", "Upah", "Komentar",
        "Status", "Detail", "Waktu Eksekusi",
    ]
    ws.append(headers)

    for item in tasks:
        idx = item["index"]
        rec = progress_data.get(str(idx), {})
        ws.append([
            idx,
            item["link"],
            item["nama_usaha"],
            item["nilai_penjualan"],
            item["biaya_pembelian"],
            item["biaya_produksi"],
            item["biaya_operasional"],
            item["upah"],
            item["komentar"],
            rec.get("status", "PENDING"),
            rec.get("detail", ""),
            rec.get("timestamp", ""),
        ])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def read_excel_source(file_path: Path) -> List[dict]:
    """Membaca daftar usaha, nilai, remark manual, dan flag timpa dari file Excel."""
    if not file_path.exists():
        console.print(f"[bold red]❌ File sumber tidak ditemukan: {file_path}[/bold red]")
        console.print("[dim]Jalankan dulu: python update_penjualan.py --init-sample[/dim]")
        sys.exit(1)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active

    headers = [
        str(c).strip().lower() if c is not None else ""
        for c in next(sheet.iter_rows(values_only=True))
    ]

    # Kolom "ekstra" (remark/catatan/timpa) tidak boleh ikut tertangkap oleh
    # pencarian kolom data utama (mis. "nilai" jangan cocok ke "remark_nilai_penjualan").
    EXTRA_PREFIXES = ("remark", "catatan", "timpa", "komentar_")

    def find_col(*keywords, fallback=None):
        for i, h in enumerate(headers):
            if h.startswith(EXTRA_PREFIXES):
                continue
            if all(k in h for k in keywords):
                return i
        return fallback

    def exact_col(*names):
        for n in names:
            if n in headers:
                return headers.index(n)
        return None

    link_col = find_col("link", fallback=0)
    nama_col = find_col("usaha", fallback=1)
    nilai_col = find_col("nilai", fallback=2)
    pembelian_col = find_col("pembelian")
    produksi_col = find_col("produksi")
    operasional_col = find_col("operasional")
    upah_col = find_col("upah")
    komentar_col = exact_col("komentar")

    remark_cols = {key: exact_col(*aliases) for key, aliases in REMARK_ALIASES.items()}
    timpa_cols = {key: exact_col(key) for key in TIMPA_KEYS}

    def cell(row, col, default=""):
        if col is None or len(row) <= col or row[col] is None:
            return default
        return _fmt_cell(row[col])

    tasks = []
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
        if not row or link_col is None or len(row) <= link_col or not row[link_col]:
            continue
        link = str(row[link_col]).strip()
        nama_usaha = cell(row, nama_col)

        if not nama_usaha:
            console.print(f"[dim yellow]↳ Baris #{row_idx} dilewati (nama_usaha kosong)[/dim yellow]")
            continue

        task = {
            "index": row_idx,
            "link": link,
            "nama_usaha": nama_usaha,
            "nilai_penjualan": cell(row, nilai_col),
            "biaya_pembelian": cell(row, pembelian_col),
            "biaya_produksi": cell(row, produksi_col),
            "biaya_operasional": cell(row, operasional_col),
            "upah": cell(row, upah_col),
            "komentar": cell(row, komentar_col) or "#1",
        }
        for key, c in remark_cols.items():
            task[key] = cell(row, c)
        for key, c in timpa_cols.items():
            task[key] = cell(row, c)
        tasks.append(task)

    return tasks


def build_remarks(item: dict) -> dict:
    """Kumpulkan remark manual + flag timpa dari satu task menjadi satu dict."""
    return {k: item.get(k) or "" for k in REMARK_KEYS + TIMPA_KEYS}


# ─────────────────────────────────────────────────────────────────────────────
# Login SSO
# ─────────────────────────────────────────────────────────────────────────────

def is_page_logged_out(page: Page) -> bool:
    """Mengecek apakah browser sedang berada pada status logout atau SSO."""
    try:
        url = page.url.lower()
        if "sso.bps.go.id" in url or "/login" in url or "/auth/login" in url or "/oauth2/authorization" in url:
            return True
        if page.locator("#kc-form-login, #username, #password, .login-button").count() > 0:
            return True
        text = page.locator("body").inner_text() if page.locator("body").count() > 0 else ""
        if "Selamat Datang Kembali" in text or "Lanjutkan dengan SSO" in text or "Login SSO BPS" in text:
            return True
        return False
    except Exception:
        return False


def perform_sso_login(page: Page, username: str = config.SSO_USERNAME, password: str = config.SSO_PASSWORD) -> bool:
    """Melakukan login otomatis SSO BPS."""
    console.print(f"[bold yellow]🔐 Login otomatis SSO BPS (User: {username})...[/bold yellow]")
    try:
        if "sso.bps.go.id" not in page.url:
            sso_btn = page.locator(
                "a:has-text('Lanjutkan dengan SSO'), a:has-text('Login SSO BPS'), a[href*='oauth2/authorization'], .login-button"
            )
            if sso_btn.count() > 0 and sso_btn.first.is_visible():
                sso_btn.first.click()
            else:
                page.goto("https://fasih-sm.bps.go.id/oauth2/authorization/ics", wait_until="domcontentloaded")
            page.wait_for_url("**/sso.bps.go.id/**", timeout=20000)

        page.wait_for_selector("#username", timeout=15000)
        page.fill("#username", username)
        page.fill("#password", password)
        page.click("#kc-login")

        page.wait_for_url("**/fasih-sm.bps.go.id/app/**", timeout=25000)
        console.print(f"[bold green]✔ Berhasil login SSO BPS sebagai {username}![/bold green]")
        time.sleep(2.0)
        return True
    except Exception as e:
        console.print(f"[bold red]❌ Gagal login otomatis SSO: {e}[/bold red]")
        return False


# ─────────────────────────────────────────────────────────────────────────────
# Navigasi & scroll
# ─────────────────────────────────────────────────────────────────────────────

def _click_menuitem(page: Page, keyword: str) -> bool:
    """Mengklik item menu di sidebar (aside) berdasarkan keyword teks."""
    return bool(page.evaluate(
        r"""(keyword) => {
        const clean = (s) => (s || "").toLowerCase().replace(/\s+/g, ' ').trim();
        const target = clean(keyword);
        const aside = document.querySelector("aside");
        if (!aside) return false;
        const items = Array.from(aside.querySelectorAll("*")).filter(el =>
            el.children.length === 0 && el.textContent.trim().length > 0
        );
        let match = items.find(el => clean(el.textContent) === target);
        if (!match) match = items.find(el => clean(el.textContent).includes(target));
        if (!match) return false;
        let clickable = match;
        while (clickable && clickable !== aside && !clickable.onclick && !String(clickable.className).includes("cursor-pointer")) {
            clickable = clickable.parentElement;
        }
        (clickable || match).scrollIntoView({ block: "center" });
        (clickable || match).click();
        return true;
    }""",
        keyword,
    ))


def _scroll_to_bottom_until(page: Page, script_find: str, step: int = 300, max_steps: int = 40, pause: float = 0.15) -> bool:
    """
    Menggulir ke bawah bertahap sampai elemen (script_find truthy) ditemukan atau
    sudah mentok di bawah. Mengembalikan True jika elemen ditemukan.
    """
    for _ in range(max_steps):
        if page.evaluate(script_find):
            return True
        before = page.evaluate("window.scrollY")
        page.evaluate(_SCROLL_STEP_JS, step)
        time.sleep(pause)
        after = page.evaluate("window.scrollY")
        more = page.evaluate(_HAS_MORE_SCROLL_JS)
        if before == after and not more:
            break
    return bool(page.evaluate(script_find))


def _scroll_to_bottom(page: Page, step: int = 300, max_steps: int = 40, pause: float = 0.15):
    """Menggulir window/layout ke paling bawah bertahap agar konten dinamis ter-render."""
    _scroll_to_bottom_until(page, "false", step=step, max_steps=max_steps, pause=pause)


def _click_usaha_item(page: Page, nama_usaha: str) -> bool:
    """
    Membuka rincian sub-blok usaha pada section KETERANGAN USAHA/PERUSAHAAN.
    Mencari kartu nama usaha di content area (bukan aside) lalu mengkliknya.
    Jika tidak ditemukan persis, akan mencari berdasarkan kemiripan (token/fuzzy).
    """
    target_js = json.dumps(nama_usaha)

    def make_script(do_click: bool) -> str:
        click_line = (
            'let clickable = match; while (clickable && clickable !== document.body && !clickable.onclick && !String(clickable.className||"").includes("cursor-pointer")) { clickable = clickable.parentElement; } (clickable || match).scrollIntoView({ block: "center" }); (clickable || match).click(); return true;'
            if do_click else "if (match) { match.scrollIntoView({ block: 'center' }); return true; } return false;"
        )
        return f"""(() => {{
        const clean = (s) => (s || "").toLowerCase().replace(/[\[\]\(\)\.\,\-\_\|\/\\:;]/g, ' ').replace(/\s+/g, ' ').trim();
        const target = clean({target_js});
        if (!target) return false;
        const targetTokens = target.split(' ').filter(x => x.length > 2);
        const nodes = Array.from(document.querySelectorAll("body *")).filter(el => {{
            if (el.closest("aside")) return false;
            if (el.children.length !== 0) return false;
            const t = clean(el.textContent);
            if (!t || t.length < 3) return false;
            return true;
        }});
        let best = null;
        let bestScore = -1;
        for (const el of nodes) {{
            const t = clean(el.textContent);
            if (!t) continue;
            let score = 0;
            if (t === target) {{ score = 1.0; }}
            else if (t.includes(target) || target.includes(t)) {{ score = Math.max(score, 0.9); }}
            else {{
                let matchCount = 0;
                for (let i = 0; i < targetTokens.length; i++) {{
                    if (t.includes(targetTokens[i])) matchCount++;
                }}
                const tokenScore = targetTokens.length ? (matchCount / targetTokens.length) : 0;
                score = Math.max(score, tokenScore * 0.8);
                if (t.length > 0 && target.length > 0) {{
                    let shorter = t.length < target.length ? t : target;
                    let longer = t.length < target.length ? target : t;
                    if (longer.includes(shorter)) score = Math.max(score, 0.7);
                }}
            }}
            if (score > bestScore) {{
                bestScore = score;
                best = el;
            }}
        }}
        if (!best || bestScore < 0.5) return false;
        {click_line}
    }})()"""

    _scroll_to_bottom(page)
    if not _scroll_to_bottom_until(page, make_script(do_click=False)):
        # coba lagi tanpa strict scroll check
        if not bool(page.evaluate(make_script(do_click=False))):
            return False
    ok = bool(page.evaluate(make_script(do_click=True)))
    if ok:
        time.sleep(1.5)
    return ok


# ─────────────────────────────────────────────────────────────────────────────
# Submit
# ─────────────────────────────────────────────────────────────────────────────

def submit_assignment(page: Page) -> Tuple[bool, str]:
    """
    Proses submit kuesioner:
    1. Masuk section CATATAN
    2. Aktifkan switch "Tampilkan Anomali Usaha dan Keluarga" (jika belum)
    3. Pastikan switch "Anomali diselesaikan oleh admin" tidak aktif
    4. Klik Kirim utama -> cek 0 Galat -> Kirim modal -> Konfirmasi final
    """
    click_catatan_ok = page.evaluate("""(() => {
        const aside = document.querySelector("aside");
        if (!aside) return false;
        const elements = Array.from(aside.querySelectorAll("*")).filter(el => el.children.length === 0 && el.textContent.trim() === "CATATAN");
        if (elements.length === 0) return false;
        const target = elements[0];
        let clickable = target;
        while (clickable && clickable !== aside && !clickable.onclick && !String(clickable.className).includes("cursor-pointer")) {
            clickable = clickable.parentElement;
        }
        (clickable || target).scrollIntoView({ block: "center" });
        (clickable || target).click();
        return true;
    })()""")

    if not click_catatan_ok:
        return False, "Menu CATATAN tidak ditemukan di sidebar"

    time.sleep(1.5)

    page.evaluate("""(() => {
        const rows = Array.from(document.querySelectorAll("div[class*='rounded-lg'], div[class*='gap-2'], div[class*='flex'], label"));

        const anomaliRow = rows.find(r => {
            const t = (r.innerText || "").toLowerCase();
            return t.includes("tampilkan anomali") && !t.includes("admin");
        });

        if (anomaliRow) {
            const input = anomaliRow.querySelector("input[role='switch'], input[type='checkbox']");
            if (input) {
                const isChecked = input.checked ||
                                  input.getAttribute("aria-checked") === "true" ||
                                  input.hasAttribute("data-checked");
                if (!isChecked) {
                    const control = anomaliRow.querySelector("[id*='control']") || input;
                    control.click();
                }
            }
        }

        const adminRow = rows.find(r => {
            const t = (r.innerText || "").toLowerCase();
            return t.includes("anomali diselesaikan oleh admin");
        });
        if (adminRow) {
            const adminInput = adminRow.querySelector("input[role='switch'], input[type='checkbox']");
            if (adminInput) {
                const isAdminChecked = adminInput.checked ||
                                       adminInput.getAttribute("aria-checked") === "true" ||
                                       adminInput.hasAttribute("data-checked");
                if (isAdminChecked) {
                    const adminControl = adminRow.querySelector("[id*='control']") || adminInput;
                    adminControl.click();
                }
            }
        }
    })()""")

    time.sleep(1.0)

    modal_opened = False
    for _ in range(3):
        page.evaluate("""(() => {
            const btns = Array.from(document.querySelectorAll("button")).filter(b => b.innerText.trim() === "Kirim" && b.offsetWidth > 0);
            const btn = btns[btns.length - 1];
            if (btn) {
                btn.scrollIntoView({ block: "center" });
                btn.click();
            }
        })()""")
        try:
            page.wait_for_selector("[class*='max-w-lg'], [role='dialog']", timeout=6000)
            modal_opened = True
            break
        except Exception:
            time.sleep(1.0)

    if not modal_opened:
        return False, "Modal ringkasan validasi tidak muncul"

    time.sleep(1.0)

    val_res = page.evaluate("""(() => {
        const modal = document.querySelector("[class*='max-w-lg'], [role='dialog']");
        if (!modal) return { status: "ERR", msg: "Modal validasi tidak ditemukan" };
        const modalText = modal.innerText;
        const galatMatch = modalText.match(/GALAT\\s*\\n\\s*(\\d+)/i);
        const galatCount = galatMatch ? parseInt(galatMatch[1]) : 0;
        if (galatCount > 0) {
            const closeBtn = modal.querySelector("button[aria-label='Dismiss']") ||
                             Array.from(modal.querySelectorAll("button")).find(b => b.innerText.trim() === "Batal");
            if (closeBtn) closeBtn.click();
            return { status: "GALAT", count: galatCount, msg: galatCount + " galat ditemukan!" };
        }
        const btns = Array.from(modal.querySelectorAll("button")).filter(b => b.innerText.trim() === "Kirim");
        if (btns.length === 0) return { status: "ERR", msg: "Tombol Kirim di modal validasi tidak ditemukan" };
        btns[0].click();
        return { status: "OK" };
    })()""")

    if val_res.get("status") != "OK":
        return False, f"Validasi gagal: {val_res.get('msg')}"

    try:
        page.wait_for_selector("button:has-text('Konfirmasi')", timeout=10000)
        page.locator("button:has-text('Konfirmasi')").first.click()
    except Exception:
        return False, "Tombol 'Konfirmasi' dialog akhir tidak ditemukan"

    try:
        page.wait_for_function("""() => {
            const t = document.body ? document.body.innerText : "";
            return t.includes("Assignment Berhasil Disubmit") || t.includes("Perubahan Anda berhasil disimpan");
        }""", timeout=15000)
        return True, "Assignment Berhasil Disubmit"
    except Exception:
        return False, "Status sukses submit tidak terdeteksi pada layar"


# ─────────────────────────────────────────────────────────────────────────────
# Wrapper, catatan, dan pengisian field
# ─────────────────────────────────────────────────────────────────────────────

def _find_wrapper_js(prefix: str, need_input: bool = True) -> str:
    """
    JS untuk mencari wrapper React berdasarkan PREFIX id (suffix dinamis, mis. '#1',
    '#1101', dan varian '_bln'). Mengembalikan {wrapper, input} atau null.
    """
    prefix_js = json.dumps(prefix)
    input_req = "true" if need_input else "false"
    return f"""(() => {{
        const els = Array.from(document.querySelectorAll('[id^=' + JSON.stringify({prefix_js}) + ']'));
        const withInput = els.filter(w => w.querySelector("input[id*='textfield-cl-']"));
        let chosen = null;
        if ({input_req}) {{
            chosen = withInput.find(w => {{
                const i = w.querySelector("input[id*='textfield-cl-']");
                return i && i.offsetWidth > 0;
            }}) || withInput[0] || null;
        }} else {{
            chosen = els.find(w => w.offsetWidth > 0) || els[0] || null;
        }}
        if (!chosen) return null;
        const inp = chosen.querySelector("input[id*='textfield-cl-']");
        return {{ wrapper: chosen.id, input: inp ? inp.id : null }};
    }})()"""


def _find_wrapper(page: Page, prefix: str, need_input: bool = True):
    """Cari wrapper via prefix; scroll ke bawah bila belum ketemu."""
    js = _find_wrapper_js(prefix, need_input)
    res = page.evaluate(js)
    if not res:
        _scroll_to_bottom(page)
        res = page.evaluate(js)
    return res


def _expand_pengeluaran(page: Page) -> bool:
    """Klik header section Pengeluaran agar field di dalamnya (mis. upah/gaji) terlihat."""
    try:
        return bool(page.evaluate("""(() => {
            const h = document.querySelector('[id^="head_pengeluaran"]');
            if (!h) return false;
            const target = h.querySelector('div') || h;
            target.scrollIntoView({ block: 'center' });
            target.click();
            return true;
        })()"""))
    except Exception:
        return False


def _close_dialog(page: Page):
    """Tutup dialog catatan (tunggu hilang, kalau tidak klik tombol close)."""
    try:
        page.wait_for_selector("[role='dialog']", state="hidden", timeout=3000)
    except Exception:
        page.evaluate(_CLOSE_DIALOG_JS)
        time.sleep(0.4)


def _write_catatan(page: Page, wrapper_exact: str, kom_to_fill: str, label: str, timpa: bool) -> Tuple[bool, str]:
    """
    Membuka dialog Catatan pada wrapper tertentu lalu:
      - membaca isi catatan lama langsung dari textarea,
      - jika sudah ada #0/#1 dan timpa=False  -> biarkan (tutup dialog),
      - jika sama persis dengan yang akan diisi -> biarkan,
      - selain itu isi/timpa dengan kom_to_fill lalu klik Simpan.
    Return (ok, status) dengan status: 'SAVED' | 'KEPT' | 'SAME' | pesan error.
    """
    wid = json.dumps(wrapper_exact)

    visible_btn_js = f"""(() => {{
        const w = document.getElementById({wid});
        if (!w) return false;
        return Array.from(w.querySelectorAll("button[title='Catatan'], button[aria-label='Catatan']"))
            .some(b => b.offsetWidth > 0 && b.offsetHeight > 0);
    }})()"""
    click_btn_js = f"""(() => {{
        const w = document.getElementById({wid});
        if (!w) return false;
        const btns = Array.from(w.querySelectorAll("button[title='Catatan'], button[aria-label='Catatan']"));
        const b = btns.find(x => x.offsetWidth > 0 && x.offsetHeight > 0) || btns[0];
        if (!b) return false;
        b.scrollIntoView({{ block: "center" }});
        b.click();
        return true;
    }})()"""

    _scroll_to_bottom_until(page, visible_btn_js)
    if not page.evaluate(click_btn_js):
        return False, f"Tombol Catatan '{label}' tidak ditemukan"

    try:
        page.wait_for_selector("textarea[placeholder*='catatan' i]", timeout=8000)
    except Exception:
        return False, f"Dialog catatan '{label}' tidak muncul"

    existing = (page.evaluate("""(() => {
        const ta = document.querySelector("textarea[placeholder*='catatan' i]");
        return ta ? (ta.value || "") : "";
    })()""") or "").strip()

    has_flag = "#0" in existing or "#1" in existing

    if existing == kom_to_fill:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return True, "SAME"

    if has_flag and not timpa:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return True, f"KEPT:{existing}"

    # Isi / timpa catatan
    res = page.evaluate(
        """(val) => {
            const ta = document.querySelector("textarea[placeholder*='catatan' i]");
            if (!ta) return "ERR";
            ta.focus();
            const setter = Object.getOwnPropertyDescriptor(window.HTMLTextAreaElement.prototype, "value").set;
            setter.call(ta, val);
            ta.dispatchEvent(new Event("input", { bubbles: true }));
            ta.dispatchEvent(new Event("change", { bubbles: true }));
            return ta.value;
        }""",
        kom_to_fill,
    )
    if res != kom_to_fill:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return False, f"Gagal mengisi textarea catatan '{label}'"

    time.sleep(0.3)
    saved = page.evaluate("""(() => {
        const b = document.querySelector("button[title='Simpan']") ||
                  Array.from(document.querySelectorAll("button")).find(x => x.innerText.trim() === "Simpan");
        if (!b) return false;
        b.click();
        return true;
    })()""")
    if not saved:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return False, f"Tombol Simpan catatan '{label}' tidak ditemukan"

    time.sleep(0.6)
    # Tutup dialog bila masih terbuka setelah Simpan
    page.evaluate(_CLOSE_DIALOG_JS)
    _close_dialog(page)
    time.sleep(0.4)
    return True, "SAVED"


def _set_nilai_field(
    page: Page,
    wrapper_prefix: str,
    value: str,
    label: str,
    komentar_manual: str = "",
    timpa: bool = False,
) -> Tuple[bool, str]:
    """
    Mengisi satu field nilai via PREFIX wrapper + catatannya.
    Komentar:
      - komentar_manual (kolom remark_* di Excel) dipakai bila terisi,
      - jika kosong: '#1' bila nilai berubah, '#0' bila sama/kosong.
    Catatan lama (#0/#1) dipertahankan kecuali timpa=True.
    """
    if value is None or str(value).strip() == "":
        console.print(f"   [dim]↳ {label}: dilewati (nilai kosong)[/dim]")
        return True, "dilewati"

    # 1) Cari wrapper + input via prefix
    found = _find_wrapper(page, wrapper_prefix, need_input=True)
    if not found or not found.get("input"):
        _expand_pengeluaran(page)  # upah/gaji kadang di section Pengeluaran yang tertutup
        time.sleep(0.8)
        found = _find_wrapper(page, wrapper_prefix, need_input=True)
    if not found or not found.get("input"):
        return False, f"TextField '{label}' (prefix {wrapper_prefix}) tidak ditemukan"
    wrapper_exact = found["wrapper"]
    input_id = found["input"]

    # 2) Isi nilai
    try:
        page.evaluate(
            "(iid) => { const e = document.getElementById(iid); if (e) e.scrollIntoView({ block: 'center' }); }",
            input_id,
        )
        time.sleep(0.6)
        inp = page.locator(f'[id="{input_id}"]')
        old_value = inp.input_value(timeout=5000)
        inp.fill(str(value), force=True, timeout=15000)
        inp.evaluate(
            "el => { el.dispatchEvent(new Event('input', {bubbles:true})); "
            "el.dispatchEvent(new Event('change', {bubbles:true})); el.blur(); }"
        )
    except Exception as e:
        return False, f"Gagal mengisi '{label}': {e}"

    console.print(f"   [dim]↳ {label}: '{old_value or '-'}' -> '{value}' (field {input_id}, wrapper {wrapper_exact})[/dim]")
    time.sleep(0.6)

    # 3) Tentukan komentar
    manual = str(komentar_manual).strip() if komentar_manual is not None else ""
    if manual:
        kom_to_fill = manual
        src = "manual dari Excel"
    else:
        old_norm = _normalize_num(old_value)
        new_norm = _normalize_num(value)
        kom_to_fill = "#0" if (old_norm == new_norm or old_norm == "" or new_norm == "") else "#1"
        src = f"otomatis, asal='{old_value}', baru='{value}'"

    # 4) Tulis catatan (baca catatan lama dari dialog, hormati flag timpa)
    ok, status = _write_catatan(page, wrapper_exact, kom_to_fill, label, timpa)
    if not ok:
        console.print(f"   [dim yellow]⚠ {status}, lanjut tanpa komentar[/dim yellow]")
        return True, "ok"

    if status == "SAVED":
        tag = "ditimpa" if timpa else "disimpan"
        console.print(f"   [dim]↳ Komentar {label}: {kom_to_fill} ({src}) → {tag}[/dim]")
    elif status == "SAME":
        console.print(f"   [dim]↳ Komentar {label}: {kom_to_fill} (sudah sama, tidak diubah)[/dim]")
    elif status.startswith("KEPT"):
        console.print(
            f"   [dim]↳ Komentar {label}: catatan lama '{status[5:]}' dipertahankan "
            f"(timpa_remark bukan 1)[/dim]"
        )
    return True, "ok"


def _set_catatan_only(
    page: Page,
    wrapper_prefix: str,
    komentar_value: str,
    label: str = "Catatan",
    timpa: bool = True,
) -> Tuple[bool, str]:
    """
    Hanya mengisi catatan pada wrapper tertentu (tanpa mengubah nilai input).
    Menerima PREFIX id (mis. 'kbli#'). Default '#0' karena nilai tidak diubah.
    """
    kom_to_fill = str(komentar_value or "").strip()
    if kom_to_fill not in ("#0", "#1"):
        kom_to_fill = "#0"

    found = _find_wrapper(page, wrapper_prefix, need_input=False)
    if not found:
        console.print(f"   [dim yellow]⚠ Wrapper '{label}' (prefix {wrapper_prefix}) tidak ditemukan[/dim yellow]")
        return False, f"Wrapper {label} tidak ditemukan"
    wrapper_exact = found["wrapper"]

    ok, status = _write_catatan(page, wrapper_exact, kom_to_fill, label, timpa)
    if not ok:
        console.print(f"   [dim yellow]⚠ {status}[/dim yellow]")
        return False, status

    console.print(f"   [dim]↳ Catatan {label} ({wrapper_exact}): '{kom_to_fill}' [{status}][/dim]")
    time.sleep(0.4)
    return True, "ok"


# ─────────────────────────────────────────────────────────────────────────────
# Alur per baris
# ─────────────────────────────────────────────────────────────────────────────

def update_penjualan_row(
    page: Page,
    link: str,
    nama_usaha: str,
    nilai_penjualan: str,
    biaya_pembelian: str = "",
    biaya_produksi: str = "",
    biaya_operasional: str = "",
    upah: str = "",
    komentar_value: str = "#1",
    kbli_catatan: str = "#0",
    kbli_wrapper_id: str = "kbli#",
    retry_on_logout: bool = True,
    remarks: Optional[dict] = None,
) -> Tuple[bool, str]:
    """
    Alur update 5 variabel nilai & komentar untuk satu usaha:
    1. Buka kuesioner (mode edit).
    2. Masuk section KETERANGAN USAHA/PERUSAHAAN, pilih sub-blok usaha.
    3. Isi tiap variabel + catatan (komentar manual / otomatis, timpa sesuai flag).
    4. Tambah catatan KBLI.
    5. Submit assignment.

    remarks: dict berisi remark_* (komentar manual) dan timpa_* (flag timpa).
    """
    if remarks is None:
        remarks = {}

    edit_url = link.strip()
    if not edit_url.endswith("/edit"):
        edit_url = edit_url.rstrip("/") + "/edit"

    values = {
        "nilai_penjualan": nilai_penjualan,
        "biaya_pembelian": biaya_pembelian,
        "biaya_produksi": biaya_produksi,
        "biaya_operasional": biaya_operasional,
        "upah": upah,
    }

    def _retry():
        return update_penjualan_row(
            page, link, nama_usaha, nilai_penjualan, biaya_pembelian, biaya_produksi,
            biaya_operasional, upah, komentar_value, kbli_catatan, kbli_wrapper_id,
            retry_on_logout=False, remarks=remarks,
        )

    try:
        page.goto(edit_url, wait_until="domcontentloaded", timeout=30000)
    except Exception as e:
        return False, f"Timeout membuka URL: {e}"

    time.sleep(1.0)

    if is_page_logged_out(page):
        if retry_on_logout and perform_sso_login(page):
            return _retry()
        return False, "Sesi login habis / terlogout"

    try:
        page.wait_for_selector("aside", timeout=20000)
    except Exception:
        if is_page_logged_out(page) and retry_on_logout:
            if perform_sso_login(page):
                return _retry()
        return False, "Menu sidebar (aside) kuesioner tidak muncul"

    time.sleep(1.0)

    # 1. Section "KETERANGAN USAHA/PERUSAHAAN"
    if not _click_menuitem(page, "keterangan usaha/perusahaan"):
        if not _click_menuitem(page, "keterangan usaha"):
            return False, "Section 'KETERANGAN USAHA/PERUSAHAAN' tidak ditemukan di sidebar"
    time.sleep(1.5)

    # 2. Sub-blok usaha
    if not _click_usaha_item(page, nama_usaha):
        return False, f"Sub-blok usaha '{nama_usaha}' tidak ditemukan"
    time.sleep(1.5)

    # 3. Isi 5 variabel (urut) + catatan
    timpa_global = _is_true(remarks.get("timpa_remark"))
    for col_key, wrapper_prefix, label, remark_key in VARIABEL_MAP:
        per_var = remarks.get("timpa_" + remark_key, "")
        # Prioritas: kolom per variabel (jika terisi, 0 pun dihormati) > kolom global
        if str(per_var).strip() != "":
            timpa = _is_true(per_var)
        else:
            timpa = timpa_global
        ok, detail = _set_nilai_field(
            page, wrapper_prefix, values.get(col_key, ""), label,
            komentar_manual=remarks.get(remark_key, ""), timpa=timpa,
        )
        if not ok:
            return False, detail

    # 4. Catatan KBLI (fallback ke kbli_genai bila kbli# tidak ada)
    if kbli_catatan and kbli_catatan.strip():
        per_kbli = remarks.get("timpa_remark_kbli", "")
        # Default KBLI: selalu isi/timpa (perilaku lama), kecuali kolom timpa_remark_kbli = 0
        timpa_kbli = _is_true(per_kbli) if str(per_kbli).strip() != "" else True

        _scroll_to_bottom(page)
        time.sleep(0.5)
        ok_kbli, det_kbli = _set_catatan_only(page, kbli_wrapper_id, kbli_catatan, "KBLI", timpa_kbli)
        if not ok_kbli:
            ok_kbli, det_kbli = _set_catatan_only(page, "kbli_genai#", kbli_catatan, "KBLI", timpa_kbli)
        if not ok_kbli:
            console.print(f"   [dim yellow]⚠ Catatan KBLI gagal: {det_kbli}[/dim yellow]")

    # 5. Submit
    console.print("   [bold cyan]📤 Melakukan Submit Assignment (CATATAN -> Kirim -> Konfirmasi)...[/bold cyan]")
    submit_ok, submit_msg = submit_assignment(page)
    if not submit_ok:
        return False, f"Gagal submit assignment: {submit_msg}"

    console.print(f"   [bold green]✔ Sukses Submit Assignment: {submit_msg}[/bold green]")
    return True, (
        f"5 Variabel + catatan (komentar otomatis #1 jika berubah, #0 jika sama; "
        f"KBLI '{kbli_catatan}') disimpan, {submit_msg}"
    )


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Batch Automation Update 5 Variabel Nilai (Blok 27) FASIH-SM SE2026 (1-Command)")
    parser.add_argument("--source", default=str(DEFAULT_SOURCE_FILE), help="File Excel sumber data (default: source/penjualan.xlsx)")
    parser.add_argument("--init-sample", action="store_true", help="Buat file Excel sumber contoh lalu keluar")
    parser.add_argument("--start", type=int, default=1, help="Nomor baris awal untuk mulai diproses (default: 1)")
    parser.add_argument("--end", type=int, default=None, help="Nomor baris akhir untuk diproses")
    parser.add_argument("--limit", type=int, default=None, help="Jumlah baris maksimum yang diproses")
    parser.add_argument("--delay", type=float, default=1.5, help="Jeda waktu antar baris dalam detik (default: 1.5)")
    parser.add_argument("--force", action="store_true", help="Paksa proses ulang semua baris meskipun sudah SUCCESS sebelumnya")
    parser.add_argument("--headless", action="store_true", help="Jalankan di latar belakang tanpa membuka jendela visual Chrome")
    parser.add_argument("--dry-run", action="store_true", help="Hanya tampilkan daftar baris tanpa mengeksekusi")
    args = parser.parse_args()

    source_path = Path(args.source)

    if args.init_sample:
        create_sample_source(source_path)
        return

    all_tasks = read_excel_source(source_path)
    total_in_file = len(all_tasks)

    progress_data = load_progress()

    completed_indices = {int(k) for k, v in progress_data.items() if v.get("status") == "SUCCESS"}
    failed_indices = {int(k) for k, v in progress_data.items() if v.get("status") == "FAILED"}

    filtered_tasks = []
    skipped_count = 0
    for t in all_tasks:
        idx = t["index"]
        if idx < args.start:
            continue
        if args.end and idx > args.end:
            continue
        if not args.force and idx in completed_indices:
            skipped_count += 1
            continue
        filtered_tasks.append(t)
        if args.limit and len(filtered_tasks) >= args.limit:
            break

    console.print(
        Panel.fit(
            f"[bold cyan]⚡ FASIH-SM Batch Update 5 Variabel Nilai (Blok 27) Engine[/bold cyan]\n"
            f"File Sumber      : [yellow]{source_path}[/yellow] (Total {total_in_file} baris)\n"
            f"SSO User         : [bold green]{config.SSO_USERNAME}[/bold green] (Auto-Login: {config.AUTO_LOGIN_SSO})\n"
            f"Profil Browser   : [magenta]{config.BROWSER_DATA_DIR}[/magenta]\n"
            f"Sudah Sukses     : [bold green]{len(completed_indices)} baris[/bold green] (Otomatis dilewati: {skipped_count})\n"
            f"Pernah Gagal     : [bold red]{len(failed_indices)} baris[/bold red] (Akan dicoba ulang)\n"
            f"Target Proses Ini: [bold cyan]{len(filtered_tasks)} baris[/bold cyan] (Mulai baris {args.start})\n"
            f"Tampilan Visual  : [yellow]{'Headless (Latar Belakang)' if args.headless else 'Google Chrome Jendela Visual'}[/yellow] | Delay: [cyan]{args.delay}s[/cyan]",
            border_style="cyan",
        )
    )

    if not filtered_tasks:
        console.print(
            "[bold green]✔ Semua baris target sudah SUCCESS sebelumnya! Tidak ada yang perlu diproses ulang.[/bold green]\n"
            "[dim]Gunakan opsi --force jika ingin tetap memproses ulang yang sudah sukses.[/dim]"
        )
        return

    if args.dry_run:
        table = Table(title="Daftar Baris Target (Dry Run)")
        table.add_column("No", justify="right", style="cyan")
        table.add_column("Nama Usaha", style="magenta")
        table.add_column("Penjualan", style="green")
        table.add_column("Pembelian", style="green")
        table.add_column("Produksi", style="green")
        table.add_column("Operasional", style="green")
        table.add_column("Upah", style="green")
        table.add_column("Timpa", style="red")
        table.add_column("Status Terakhir", style="yellow")

        for t in filtered_tasks[:20]:
            prev_status = progress_data.get(str(t["index"]), {}).get("status", "PENDING")
            timpa_info = []
            if _is_true(t.get("timpa_remark")):
                timpa_info.append("semua")
            for k in REMARK_KEYS:
                if str(t.get("timpa_" + k, "")).strip() != "":
                    timpa_info.append(f"{k.replace('remark_', '')}={t['timpa_' + k]}")
            table.add_row(
                str(t["index"]), t["nama_usaha"],
                t["nilai_penjualan"], t["biaya_pembelian"],
                t["biaya_produksi"], t["biaya_operasional"], t["upah"],
                ", ".join(timpa_info) or "-", prev_status,
            )

        console.print(table)
        if len(filtered_tasks) > 20:
            console.print(f"[dim]... dan {len(filtered_tasks) - 20} baris lainnya.[/dim]")
        return

    console.print(f"[cyan]🚀 Meluncurkan Google Chrome otomatis ({'visual' if not args.headless else 'headless'})...[/cyan]")
    config.BROWSER_DATA_DIR.mkdir(parents=True, exist_ok=True)

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(config.BROWSER_DATA_DIR),
            channel=config.BROWSER_CHANNEL,
            headless=args.headless,
            viewport=None,
            args=[
                "--start-maximized",
                "--disable-blink-features=AutomationControlled",
            ],
        )

        page = context.pages[0] if context.pages else context.new_page()

        try:
            page.goto(f"{config.FASIH_BASE_URL}/app/", wait_until="domcontentloaded", timeout=25000)
            time.sleep(2.0)
            if is_page_logged_out(page):
                console.print("[yellow]⚠ Status awal terdeteksi logout! Melakukan login otomatis SSO BPS...[/yellow]")
                if not perform_sso_login(page):
                    console.print("[bold red]❌ Gagal login otomatis SSO. Proses dibatalkan.[/bold red]")
                    context.close()
                    return
        except Exception as e:
            console.print(f"[dim yellow]Peringatan cek awal: {e}[/dim yellow]")

        success_count = 0
        fail_count = 0
        current_item = None

        try:
            for i, item in enumerate(filtered_tasks, start=1):
                current_item = item
                idx = item["index"]
                link = item["link"]
                nama_usaha = item["nama_usaha"]
                nilai = item["nilai_penjualan"]
                pembelian = item["biaya_pembelian"]
                produksi = item["biaya_produksi"]
                operasional = item["biaya_operasional"]
                upah = item["upah"]
                komentar = item["komentar"]
                curr_time_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

                console.print(f"\n[bold blue]━━━ [{i}/{len(filtered_tasks)}] Memproses Baris #{idx} ━━━[/bold blue]")
                console.print(f"Usaha   : [magenta]{nama_usaha}[/magenta]")
                console.print(
                    f"Nilai   : Penjualan=[cyan]{nilai or '-'}[/cyan] "
                    f"Pembelian=[cyan]{pembelian or '-'}[/cyan] "
                    f"Produksi=[cyan]{produksi or '-'}[/cyan] "
                    f"Operasional=[cyan]{operasional or '-'}[/cyan] "
                    f"Upah=[cyan]{upah or '-'}[/cyan] | Komentar: [cyan]{komentar}[/cyan]"
                )
                console.print(f"Link    : [dim]{link}[/dim]")

                remarks = build_remarks(item)
                try:
                    is_ok, msg = update_penjualan_row(
                        page, link, nama_usaha, nilai, pembelian, produksi,
                        operasional, upah, komentar, remarks=remarks,
                    )
                except KeyboardInterrupt:
                    raise
                except Exception as e:
                    is_ok, msg = False, f"Exception tak terduga: {e}"

                status_str = "SUCCESS" if is_ok else "FAILED"
                progress_data[str(idx)] = {
                    "link": link,
                    "nama_usaha": nama_usaha,
                    "nilai_penjualan": nilai,
                    "biaya_pembelian": pembelian,
                    "biaya_produksi": produksi,
                    "biaya_operasional": operasional,
                    "upah": upah,
                    "komentar": komentar,
                    "status": status_str,
                    "detail": msg,
                    "timestamp": curr_time_str,
                }
                save_progress(progress_data)
                current_item = None

                if is_ok:
                    success_count += 1
                    console.print(f"[bold green]✔ #{idx} Berhasil: {msg}[/bold green]")
                else:
                    fail_count += 1
                    console.print(f"[bold red]❌ #{idx} Gagal: {msg}[/bold red]")

                if i % 10 == 0 or i == len(filtered_tasks):
                    export_result_excel(all_tasks, progress_data, RESULT_EXCEL_FILE)

                if i < len(filtered_tasks) and args.delay > 0:
                    time.sleep(args.delay)

        except KeyboardInterrupt:
            console.print("\n[bold yellow]⚠ Eksekusi dihentikan pengguna (Ctrl+C). Checkpoint tersimpan aman![/bold yellow]")
            if current_item is not None:
                idx_int = current_item["index"]
                progress_data[str(idx_int)] = {
                    "link": current_item["link"],
                    "nama_usaha": current_item["nama_usaha"],
                    "nilai_penjualan": current_item["nilai_penjualan"],
                    "biaya_pembelian": current_item["biaya_pembelian"],
                    "biaya_produksi": current_item["biaya_produksi"],
                    "biaya_operasional": current_item["biaya_operasional"],
                    "upah": current_item["upah"],
                    "komentar": current_item["komentar"],
                    "status": "FAILED",
                    "detail": "Dihentikan pengguna (Ctrl+C) saat sedang diproses",
                    "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
                save_progress(progress_data)
                console.print(f"[yellow]↳ Baris #{idx_int} ditandai FAILED (terhenti) & akan dicoba ulang.[/yellow]")
        finally:
            export_result_excel(all_tasks, progress_data, RESULT_EXCEL_FILE)
            context.close()
            console.print(
                Panel.fit(
                    f"[bold cyan]📊 Rekapitulasi Update 5 Variabel Nilai[/bold cyan]\n"
                    f"Selesai Diproses : [bold]{success_count + fail_count}[/bold] baris\n"
                    f"Berhasil (Submit): [bold green]{success_count}[/bold green]\n"
                    f"Gagal            : [bold red]{fail_count}[/bold red]\n"
                    f"Log Checkpoint   : [yellow]{PROGRESS_FILE}[/yellow]\n"
                    f"Hasil Excel      : [green]{RESULT_EXCEL_FILE}[/green]",
                    border_style="cyan",
                )
            )


if __name__ == "__main__":
    main()