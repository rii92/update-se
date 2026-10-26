#!/usr/bin/env python3
"""
Batch Automation Remark NIK FASIH-SM SE2026
Menambahkan catatan (remark) pada field NIK anggota keluarga di section
"KETERANGAN UMUM KELUARGA" berdasarkan file Excel (source/remark_nik.xlsx)
hanya dengan 1 perintah:

1. Otomatis membuka Google Chrome dengan profil terisolasi (config.BROWSER_DATA_DIR).
2. Otomatis login SSO BPS jika belum login / sesi habis.
3. Otomatis melewati baris yang sudah SUCCESS sebelumnya (checkpoint).
4. Untuk setiap baris:
   a. Membuka link kuesioner dalam mode edit (.../edit).
   b. Masuk section "KETERANGAN UMUM KELUARGA".
   c. Klik anggota keluarga sesuai kolom nama.
   d. Klik tombol Catatan pada field NIK (wrapper prefix: nik_dtsen#).
   e. Mengisi catatan (default "#0"; catatan lama dipertahankan kecuali
      timpa_remark = 1), simpan lalu tutup dialog.
   f. Submit assignment (masuk CATATAN -> switch anomali -> Kirim -> Konfirmasi).
5. Otomatis merekap hasil ke source/remark_nik_result.xlsx.

Kolom file sumber (source/remark_nik.xlsx):
    link          : link assignment FASIH (boleh tanpa /edit)
    nama_art      : nama anggota keluarga persis seperti di daftar (atribut title)
    nik           : (opsional) hanya untuk tampilan log / rekap
    komentar      : (opsional) isi catatan, default "#0"
    timpa_remark  : (opsional) 1 = timpa catatan lama yang sudah ada (#0/#1);
                    kosong/0 = catatan lama dipertahankan

Contoh:
    python remark_nik.py --init-sample
    python remark_nik.py
    python remark_nik.py --headless
    python remark_nik.py --dry-run --limit 10
    python remark_nik.py --start 1 --end 50
    python remark_nik.py --force
"""

import argparse
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from playwright.sync_api import Page, sync_playwright
from rich.console import Console
from rich.panel import Panel
from rich.table import Table

import config

console = Console()

PROGRESS_FILE = Path("source/remark_nik_progress.json")
RESULT_EXCEL_FILE = Path("source/remark_nik_result.xlsx")
DEFAULT_SOURCE_FILE = Path("source/remark_nik.xlsx")

NIK_WRAPPER_PREFIX = "nik_dtsen#"          # id wrapper field NIK (suffix dinamis: #3, #4, ...)
SECTION_KEYWORD = "keterangan umum keluarga"
DEFAULT_KOMENTAR = "#0"
KIRIM_DELAY = 10.0                         # jeda (detik) sebelum klik Kirim; ubah via --delay-kirim

_SCROLL_STEP_JS = """(step) => {
    window.scrollBy({ top: step, behavior: "instant" });
    const scrollers = Array.from(document.querySelectorAll("div")).filter(el => {
        const oy = getComputedStyle(el).overflowY;
        return (oy === "auto" || oy === "scroll") && el.scrollHeight > el.clientHeight + 20;
    });
    scrollers.forEach(el => el.scrollBy({ top: step, behavior: "instant" }));
}"""

_HAS_MORE_SCROLL_JS = """(() => {
    const scrollers = Array.from(document.querySelectorAll("div")).filter(el => {
        const oy = getComputedStyle(el).overflowY;
        return (oy === "auto" || oy === "scroll") && el.scrollHeight > el.clientHeight + 20;
    });
    return scrollers.some(el => el.scrollTop + el.clientHeight < el.scrollHeight - 5);
})()"""

_CLOSE_DIALOG_JS = """(() => {
    const d = document.querySelector("[role='dialog']") || document.querySelector("[class*='max-w-lg']");
    const scope = d || document;
    const b = scope.querySelector("button[aria-label='Dismiss']") ||
              Array.from(scope.querySelectorAll("button")).find(x => {
                  const t = x.innerText.trim();
                  return t === "Close" || t === "Batal";
              });
    if (b) { b.click(); return true; }
    return false;
})()"""


# ─────────────────────────────────────────────────────────────────────────────
# Utilitas
# ─────────────────────────────────────────────────────────────────────────────

def _is_true(val) -> bool:
    if val is None:
        return False
    return str(val).strip().lower() in ("1", "1.0", "true", "ya", "y", "yes")


def _fmt_cell(val) -> str:
    """Rapikan nilai sel Excel. Float bulat -> int; angka panjang (NIK) tidak jadi 3.2e15."""
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


# ─────────────────────────────────────────────────────────────────────────────
# Excel sumber / progress / hasil
# ─────────────────────────────────────────────────────────────────────────────

def create_sample_source(file_path: Path = DEFAULT_SOURCE_FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Remark NIK"
    ws.append(["link", "nama_art", "nik", "komentar", "timpa_remark"])
    ws.append([config.DEFAULT_ASSIGNMENT_URL, "NANDA JAGO SAPUTRA", "'6171000000000001", "#0", "0"])
    file_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(file_path)
    console.print(f"[bold green]✔ Contoh file sumber dibuat:[/bold green] [yellow]{file_path}[/yellow]")
    console.print(
        "[dim]Kolom: link, nama_art, nik (opsional), komentar (opsional, default #0), "
        "timpa_remark (opsional, 1 = timpa catatan lama). Satu baris per anggota keluarga.[/dim]"
    )


def read_excel_source(file_path: Path) -> List[dict]:
    if not file_path.exists():
        console.print(f"[bold red]❌ File sumber tidak ditemukan: {file_path}[/bold red]")
        console.print("[dim]Jalankan dulu: python remark_nik.py --init-sample[/dim]")
        sys.exit(1)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active
    headers = [
        str(c).strip().lower() if c is not None else ""
        for c in next(sheet.iter_rows(values_only=True))
    ]

    def exact_col(*names):
        for n in names:
            if n in headers:
                return headers.index(n)
        return None

    link_col = exact_col("link")
    if link_col is None:
        link_col = next((i for i, h in enumerate(headers) if "link" in h), 0)

    nama_col = exact_col("nama_art", "nama", "nama_anggota", "nama_anggota_keluarga", "nama_krt")
    if nama_col is None:
        nama_col = next((i for i, h in enumerate(headers) if "nama" in h), 1)

    nik_col = exact_col("nik")
    komentar_col = exact_col("komentar", "remark", "remark_nik", "catatan")
    timpa_col = exact_col("timpa_remark", "timpa")

    def cell(row, col, default=""):
        if col is None or len(row) <= col or row[col] is None:
            return default
        return _fmt_cell(row[col])

    tasks = []
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
        if not row or len(row) <= link_col or not row[link_col]:
            continue
        nama = cell(row, nama_col)
        if not nama:
            console.print(f"[dim yellow]↳ Baris #{row_idx} dilewati (nama_art kosong)[/dim yellow]")
            continue
        tasks.append({
            "index": row_idx,
            "link": str(row[link_col]).strip(),
            "nama": nama,
            "nik": cell(row, nik_col).lstrip("'"),
            "komentar": cell(row, komentar_col) or DEFAULT_KOMENTAR,
            "timpa_remark": cell(row, timpa_col),
        })
    return tasks


def load_progress() -> Dict[str, dict]:
    progress: Dict[str, dict] = {}
    if PROGRESS_FILE.exists():
        try:
            with open(PROGRESS_FILE, "r", encoding="utf-8") as f:
                progress = json.load(f)
        except Exception:
            pass

    # Fallback: pulihkan status dari Excel hasil bila JSON hilang
    if RESULT_EXCEL_FILE.exists():
        try:
            wb = openpyxl.load_workbook(RESULT_EXCEL_FILE, data_only=True)
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            header = [str(h).strip().lower() if h is not None else "" for h in (next(rows, None) or [])]

            def col(name):
                return header.index(name) if name in header else None

            c_no, c_link, c_nama, c_nik = col("no"), col("link"), col("nama art"), col("nik")
            c_kom, c_status, c_detail, c_waktu = col("komentar"), col("status"), col("detail"), col("waktu eksekusi")

            def get(row, c):
                return "" if c is None or len(row) <= c or row[c] is None else str(row[c]).strip()

            for row in rows:
                no = get(row, c_no)
                status = get(row, c_status)
                if not no or not status:
                    continue
                try:
                    no = str(int(float(no)))
                except ValueError:
                    pass
                if no not in progress:
                    progress[no] = {
                        "link": get(row, c_link), "nama": get(row, c_nama), "nik": get(row, c_nik),
                        "komentar": get(row, c_kom), "status": status,
                        "detail": get(row, c_detail), "timestamp": get(row, c_waktu),
                    }
        except Exception:
            pass
    return progress


def save_progress(data: Dict[str, dict]):
    """Simpan checkpoint secara atomik."""
    PROGRESS_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = PROGRESS_FILE.with_suffix(PROGRESS_FILE.suffix + ".tmp")
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp_path, PROGRESS_FILE)


def export_result_excel(tasks: List[dict], progress_data: Dict[str, dict], output_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hasil Remark NIK"
    ws.append(["No", "Link", "Nama ART", "NIK", "Komentar", "Status", "Detail", "Waktu Eksekusi"])
    for item in tasks:
        rec = progress_data.get(str(item["index"]), {})
        ws.append([
            item["index"], item["link"], item["nama"], item["nik"], item["komentar"],
            rec.get("status", "PENDING"), rec.get("detail", ""), rec.get("timestamp", ""),
        ])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


# ─────────────────────────────────────────────────────────────────────────────
# Login SSO
# ─────────────────────────────────────────────────────────────────────────────

def is_page_logged_out(page: Page) -> bool:
    try:
        url = page.url.lower()
        if "sso.bps.go.id" in url or "/login" in url or "/auth/login" in url or "/oauth2/authorization" in url:
            return True
        if page.locator("#kc-form-login, #username, #password, .login-button").count() > 0:
            return True
        text = page.locator("body").inner_text() if page.locator("body").count() > 0 else ""
        return "Selamat Datang Kembali" in text or "Lanjutkan dengan SSO" in text or "Login SSO BPS" in text
    except Exception:
        return False


def perform_sso_login(page: Page, username: str = config.SSO_USERNAME, password: str = config.SSO_PASSWORD) -> bool:
    console.print(f"[bold yellow]🔐 Login otomatis SSO BPS (User: {username})...[/bold yellow]")
    try:
        if "sso.bps.go.id" not in page.url:
            sso_btn = page.locator(
                "a:has-text('Lanjutkan dengan SSO'), a:has-text('Login SSO BPS'), "
                "a[href*='oauth2/authorization'], .login-button"
            )
            if sso_btn.count() > 0 and sso_btn.first.is_visible():
                sso_btn.first.click()
            else:
                page.goto("https://fasih-sm.bps.go.id/oauth2/authorization/ics", wait_until="domcontentloaded")
            page.wait_for_url("**/sso.bps.go.id/**", timeout=20000)

        page.wait_for_selector("#username", timeout=15000)
        page.fill("#username", username)
        page.fill("#password", password)
        page.click("#kc-login")
        page.wait_for_url("**/fasih-sm.bps.go.id/app/**", timeout=25000)
        console.print(f"[bold green]✔ Berhasil login SSO BPS sebagai {username}![/bold green]")
        time.sleep(2.0)
        return True
    except Exception as e:
        console.print(f"[bold red]❌ Gagal login otomatis SSO: {e}[/bold red]")
        return False


# ─────────────────────────────────────────────────────────────────────────────
# Navigasi
# ─────────────────────────────────────────────────────────────────────────────

def _click_menuitem(page: Page, keyword: str) -> bool:
    """Klik item menu sidebar (aside) berdasarkan teks."""
    return bool(page.evaluate(
        r"""(keyword) => {
        const clean = (s) => (s || "").toLowerCase().replace(/\s+/g, ' ').trim();
        const target = clean(keyword);
        const aside = document.querySelector("aside");
        if (!aside) return false;
        const items = Array.from(aside.querySelectorAll("*")).filter(el =>
            el.children.length === 0 && el.textContent.trim().length > 0
        );
        let match = items.find(el => clean(el.textContent) === target);
        if (!match) match = items.find(el => clean(el.textContent).includes(target));
        if (!match) return false;
        let clickable = match;
        while (clickable && clickable !== aside && !clickable.onclick && !String(clickable.className).includes("cursor-pointer")) {
            clickable = clickable.parentElement;
        }
        (clickable || match).scrollIntoView({ block: "center" });
        (clickable || match).click();
        return true;
    }""",
        keyword,
    ))


def _scroll_to_bottom_until(page: Page, script_find: str, step: int = 300, max_steps: int = 40, pause: float = 0.15) -> bool:
    """Gulir bertahap sampai script_find truthy atau mentok. True jika ditemukan."""
    for _ in range(max_steps):
        if page.evaluate(script_find):
            return True
        before = page.evaluate("window.scrollY")
        page.evaluate(_SCROLL_STEP_JS, step)
        time.sleep(pause)
        after = page.evaluate("window.scrollY")
        more = page.evaluate(_HAS_MORE_SCROLL_JS)
        if before == after and not more:
            break
    return bool(page.evaluate(script_find))


def _scroll_to_bottom(page: Page):
    _scroll_to_bottom_until(page, "false")


def _click_member(page: Page, nama: str) -> bool:
    """
    Klik anggota keluarga berdasarkan nama (atribut title, seperti di recording:
    getByTitle(nama).first()). Fallback: pencocokan teks di content area.
    """
    esc = nama.replace("\\", "\\\\").replace('"', '\\"')

    def try_click(loc) -> bool:
        try:
            if loc.count() == 0:
                return False
            loc.first.scroll_into_view_if_needed(timeout=5000)
            loc.first.click(timeout=8000)
            return True
        except Exception:
            return False

    for _ in range(2):
        if try_click(page.locator(f'[title="{esc}"]:not(aside *)')):
            return True
        if try_click(page.locator(f'[title="{esc}"]')):
            return True
        if try_click(page.get_by_title(nama, exact=False)):
            return True
        _scroll_to_bottom(page)

    target_js = json.dumps(nama)
    script = r"""(() => {
        const clean = (s) => (s || "").toLowerCase().replace(/\s+/g, ' ').trim();
        const target = clean(%s);
        const nodes = Array.from(document.querySelectorAll("body *")).filter(el =>
            el.children.length === 0 && clean(el.textContent) === target && !el.closest("aside")
        );
        if (!nodes.length) return false;
        nodes[0].scrollIntoView({ block: "center" });
        nodes[0].click();
        return true;
    })()""" % target_js
    return bool(page.evaluate(script))


def _find_wrapper(page: Page, prefix: str, timeout: float = 10.0) -> Optional[str]:
    """Cari id wrapper via PREFIX id (suffix dinamis). Utamakan yang terlihat."""
    js = """(prefix) => {
        const els = Array.from(document.querySelectorAll("[id]")).filter(e => e.id.startsWith(prefix));
        const chosen = els.find(e => e.offsetWidth > 0 && e.offsetHeight > 0) || els[0] || null;
        return chosen ? chosen.id : null;
    }"""
    deadline = time.time() + timeout
    while time.time() < deadline:
        wid = page.evaluate(js, prefix)
        if wid:
            return wid
        page.evaluate(_SCROLL_STEP_JS, 300)
        time.sleep(0.4)
    return page.evaluate(js, prefix)


# ─────────────────────────────────────────────────────────────────────────────
# Catatan
# ─────────────────────────────────────────────────────────────────────────────

def _close_dialog(page: Page):
    try:
        page.wait_for_selector("[role='dialog']", state="hidden", timeout=3000)
    except Exception:
        page.evaluate(_CLOSE_DIALOG_JS)
        time.sleep(0.4)


def _write_catatan(page: Page, wrapper_id: str, kom_to_fill: str, label: str, timpa: bool) -> Tuple[bool, str]:
    """
    Buka dialog Catatan pada wrapper, baca catatan lama dari textarea, lalu:
      - sama persis dengan yang akan diisi  -> SAME
      - sudah ada #0/#1 dan timpa=False      -> KEPT:<isi lama>
      - selain itu isi/timpa, Simpan, tutup  -> SAVED
    """
    wid = json.dumps(wrapper_id)

    visible_btn_js = f"""(() => {{
        const w = document.getElementById({wid});
        if (!w) return false;
        return Array.from(w.querySelectorAll("button[title='Catatan'], button[aria-label='Catatan']"))
            .some(b => b.offsetWidth > 0 && b.offsetHeight > 0);
    }})()"""
    click_btn_js = f"""(() => {{
        const w = document.getElementById({wid});
        if (!w) return false;
        const btns = Array.from(w.querySelectorAll("button[title='Catatan'], button[aria-label='Catatan']"));
        const b = btns.find(x => x.offsetWidth > 0 && x.offsetHeight > 0) || btns[0];
        if (!b) return false;
        b.scrollIntoView({{ block: "center" }});
        b.click();
        return true;
    }})()"""

    page.evaluate(f"""(() => {{ const w = document.getElementById({wid}); if (w) w.scrollIntoView({{ block: "center" }}); }})()""")
    time.sleep(0.4)
    _scroll_to_bottom_until(page, visible_btn_js, max_steps=10)
    if not page.evaluate(click_btn_js):
        return False, f"Tombol Catatan '{label}' tidak ditemukan"

    try:
        page.wait_for_selector("textarea[placeholder*='catatan' i]", timeout=8000)
    except Exception:
        return False, f"Dialog catatan '{label}' tidak muncul"

    existing = (page.evaluate("""(() => {
        const ta = document.querySelector("textarea[placeholder*='catatan' i]");
        return ta ? (ta.value || "") : "";
    })()""") or "").strip()

    if existing == kom_to_fill:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return True, "SAME"

    if ("#0" in existing or "#1" in existing) and not timpa:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return True, f"KEPT:{existing}"

    res = page.evaluate(
        """(val) => {
            const ta = document.querySelector("textarea[placeholder*='catatan' i]");
            if (!ta) return "ERR";
            ta.focus();
            const setter = Object.getOwnPropertyDescriptor(window.HTMLTextAreaElement.prototype, "value").set;
            setter.call(ta, val);
            ta.dispatchEvent(new Event("input", { bubbles: true }));
            ta.dispatchEvent(new Event("change", { bubbles: true }));
            return ta.value;
        }""",
        kom_to_fill,
    )
    if res != kom_to_fill:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return False, f"Gagal mengisi textarea catatan '{label}'"

    time.sleep(0.3)
    # Tombol Simpan opsional (di recording catatan langsung ditutup dengan Dismiss)
    page.evaluate("""(() => {
        const b = document.querySelector("button[title='Simpan']") ||
                  Array.from(document.querySelectorAll("button")).find(x => x.innerText.trim() === "Simpan");
        if (b) b.click();
    })()""")
    time.sleep(0.6)
    page.evaluate(_CLOSE_DIALOG_JS)   # klik Dismiss
    _close_dialog(page)
    time.sleep(0.4)
    return True, "SAVED"


# ─────────────────────────────────────────────────────────────────────────────
# Submit
# ─────────────────────────────────────────────────────────────────────────────

def submit_assignment(page: Page) -> Tuple[bool, str]:
    """
    Proses submit kuesioner (sama seperti update_penjualan.py / cleaning_nik.py):
    1. Masuk ke section CATATAN
    2. Cek switch "Tampilkan Anomali Usaha dan Keluarga" (aktifkan jika belum)
    3. Pastikan switch "Anomali diselesaikan oleh admin" tidak aktif
    4. Klik Kirim utama -> Cek validasi 0 Galat -> Klik Kirim modal -> Konfirmasi final
    """
    # 1. Masuk ke section CATATAN
    click_catatan_ok = page.evaluate("""(() => {
        const aside = document.querySelector("aside");
        if (!aside) return false;
        const elements = Array.from(aside.querySelectorAll("*")).filter(el => el.children.length === 0 && el.textContent.trim() === "CATATAN");
        if (elements.length === 0) return false;
        const target = elements[0];
        let clickable = target;
        while (clickable && clickable !== aside && !clickable.onclick && !String(clickable.className).includes("cursor-pointer")) {
            clickable = clickable.parentElement;
        }
        (clickable || target).scrollIntoView({ block: "center" });
        (clickable || target).click();
        return true;
    })()""")

    if not click_catatan_ok:
        return False, "Menu CATATAN tidak ditemukan di sidebar"

    time.sleep(1.5)

    # 2. Cek switch di blok CATATAN:
    # - "Tampilkan Anomali Usaha dan Keluarga" jika sudah dichecklist JANGAN dichecklist lagi.
    # - JANGAN checklist "Anomali diselesaikan oleh admin".
    # - JANGAN memberi catatan penjelasan anomali admin.
    page.evaluate("""(() => {
        const rows = Array.from(document.querySelectorAll("div[class*='rounded-lg'], div[class*='gap-2'], div[class*='flex'], label"));

        // Cari switch "Tampilkan Anomali Usaha dan Keluarga" (hindari baris admin)
        const anomaliRow = rows.find(r => {
            const t = (r.innerText || "").toLowerCase();
            return t.includes("tampilkan anomali") && !t.includes("admin");
        });

        if (anomaliRow) {
            const input = anomaliRow.querySelector("input[role='switch'], input[type='checkbox']");
            if (input) {
                const isChecked = input.checked ||
                                  input.getAttribute("aria-checked") === "true" ||
                                  input.hasAttribute("data-checked");
                if (!isChecked) {
                    const control = anomaliRow.querySelector("[id*='control']") || input;
                    control.click();
                }
            }
        }

        // Pastikan switch "Anomali diselesaikan oleh admin" TIDAK DICHECKLIST (jika sempat aktif, uncheck)
        const adminRow = rows.find(r => {
            const t = (r.innerText || "").toLowerCase();
            return t.includes("anomali diselesaikan oleh admin");
        });
        if (adminRow) {
            const adminInput = adminRow.querySelector("input[role='switch'], input[type='checkbox']");
            if (adminInput) {
                const isAdminChecked = adminInput.checked ||
                                       adminInput.getAttribute("aria-checked") === "true" ||
                                       adminInput.hasAttribute("data-checked");
                if (isAdminChecked) {
                    const adminControl = adminRow.querySelector("[id*='control']") || adminInput;
                    adminControl.click();
                }
            }
        }
    })()""")

    time.sleep(1.0)

    # Jeda sebelum klik Kirim (beri waktu catatan & switch tersimpan di server)
    if KIRIM_DELAY > 0:
        console.print(f"   [dim]⏳ Menunggu {KIRIM_DELAY:g} detik sebelum klik Kirim...[/dim]")
        time.sleep(KIRIM_DELAY)

    # 3. Klik tombol Kirim utama (ulangi bila modal validasi belum muncul)
    modal_opened = False
    for _ in range(3):
        page.evaluate("""(() => {
            const btns = Array.from(document.querySelectorAll("button")).filter(b => b.innerText.trim() === "Kirim" && b.offsetWidth > 0);
            const btn = btns[btns.length - 1];
            if (btn) {
                btn.scrollIntoView({ block: "center" });
                btn.click();
            }
        })()""")
        try:
            page.wait_for_selector("[class*='max-w-lg'], [role='dialog']", timeout=6000)
            modal_opened = True
            break
        except Exception:
            time.sleep(1.0)

    if not modal_opened:
        return False, "Modal ringkasan validasi tidak muncul"

    time.sleep(1.0)

    # 4. Validasi 0 Galat & Klik Kirim di modal
    val_res = page.evaluate("""(() => {
        const modal = document.querySelector("[class*='max-w-lg'], [role='dialog']");
        if (!modal) return { status: "ERR", msg: "Modal validasi tidak ditemukan" };
        const modalText = modal.innerText;
        const galatMatch = modalText.match(/GALAT\\s*\\n\\s*(\\d+)/i);
        const galatCount = galatMatch ? parseInt(galatMatch[1]) : 0;
        if (galatCount > 0) {
            const closeBtn = modal.querySelector("button[aria-label='Dismiss']") ||
                             Array.from(modal.querySelectorAll("button")).find(b => b.innerText.trim() === "Batal");
            if (closeBtn) closeBtn.click();
            return { status: "GALAT", count: galatCount, msg: galatCount + " galat ditemukan!" };
        }
        const btns = Array.from(modal.querySelectorAll("button")).filter(b => b.innerText.trim() === "Kirim");
        if (btns.length === 0) return { status: "ERR", msg: "Tombol Kirim di modal validasi tidak ditemukan" };
        btns[0].click();
        return { status: "OK" };
    })()""")

    if val_res.get("status") != "OK":
        return False, f"Validasi gagal: {val_res.get('msg')}"

    # 5. Klik tombol Konfirmasi pada dialog akhir
    try:
        page.wait_for_selector("button:has-text('Konfirmasi')", timeout=10000)
        page.locator("button:has-text('Konfirmasi')").first.click()
    except Exception:
        return False, "Tombol 'Konfirmasi' dialog akhir tidak ditemukan"

    # 6. Verifikasi submit selesai
    try:
        page.wait_for_function("""() => {
            const t = document.body ? document.body.innerText : "";
            return t.includes("Assignment Berhasil Disubmit") || t.includes("Perubahan Anda berhasil disimpan");
        }""", timeout=15000)
        return True, "Assignment Berhasil Disubmit"
    except Exception:
        return False, "Status sukses submit tidak terdeteksi pada layar"


# ─────────────────────────────────────────────────────────────────────────────
# Alur per baris
# ─────────────────────────────────────────────────────────────────────────────

def remark_nik_row(
    page: Page,
    link: str,
    nama: str,
    komentar: str = DEFAULT_KOMENTAR,
    timpa: bool = False,
    retry_on_logout: bool = True,
) -> Tuple[bool, str]:
    edit_url = link.strip()
    if not edit_url.endswith("/edit"):
        edit_url = edit_url.rstrip("/") + "/edit"

    def _retry():
        return remark_nik_row(page, link, nama, komentar, timpa, retry_on_logout=False)

    try:
        page.goto(edit_url, wait_until="domcontentloaded", timeout=30000)
    except Exception as e:
        return False, f"Timeout membuka URL: {e}"
    time.sleep(1.0)

    if is_page_logged_out(page):
        if retry_on_logout and perform_sso_login(page):
            return _retry()
        return False, "Sesi login habis / terlogout"

    try:
        page.wait_for_selector("aside", timeout=20000)
    except Exception:
        if is_page_logged_out(page) and retry_on_logout and perform_sso_login(page):
            return _retry()
        return False, "Menu sidebar (aside) kuesioner tidak muncul"
    time.sleep(1.0)

    # 1. Section KETERANGAN UMUM KELUARGA
    if not _click_menuitem(page, SECTION_KEYWORD):
        return False, "Section 'KETERANGAN UMUM KELUARGA' tidak ditemukan di sidebar"
    time.sleep(1.5)

    # 2. Pilih anggota keluarga
    if not _click_member(page, nama):
        return False, f"Anggota keluarga '{nama}' tidak ditemukan"
    time.sleep(1.5)

    # 3. Catatan pada field NIK
    wrapper_id = _find_wrapper(page, NIK_WRAPPER_PREFIX)
    if not wrapper_id:
        return False, f"Field NIK (prefix {NIK_WRAPPER_PREFIX}) tidak ditemukan"

    ok, status = _write_catatan(page, wrapper_id, komentar, "NIK", timpa)
    if not ok:
        return False, status

    if status == "SAVED":
        console.print(f"   [dim]↳ Catatan NIK ({wrapper_id}): '{komentar}' {'(ditimpa)' if timpa else '(disimpan)'}[/dim]")
    elif status == "SAME":
        console.print(f"   [dim]↳ Catatan NIK ({wrapper_id}): '{komentar}' sudah sama, tidak diubah[/dim]")
    else:
        console.print(f"   [dim]↳ Catatan NIK: catatan lama '{status[5:]}' dipertahankan (timpa_remark bukan 1)[/dim]")

    # 4. Submit assignment (CATATAN -> switch anomali -> Kirim -> Konfirmasi)
    console.print("   [bold cyan]📤 Melakukan Submit Assignment (CATATAN -> Kirim -> Konfirmasi)...[/bold cyan]")
    submit_ok, submit_msg = submit_assignment(page)
    if not submit_ok:
        return False, f"Gagal submit assignment: {submit_msg}"

    console.print(f"   [bold green]✔ Sukses Submit Assignment: {submit_msg}[/bold green]")
    return True, f"Catatan NIK '{komentar}' [{status.split(':')[0]}], {submit_msg}"


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

def main():
    global KIRIM_DELAY
    parser = argparse.ArgumentParser(description="Batch Automation Remark NIK FASIH-SM SE2026 (1-Command)")
    parser.add_argument("--source", default=str(DEFAULT_SOURCE_FILE), help="File Excel sumber (default: source/remark_nik.xlsx)")
    parser.add_argument("--init-sample", action="store_true", help="Buat file Excel sumber contoh lalu keluar")
    parser.add_argument("--start", type=int, default=1, help="Nomor baris awal (default: 1)")
    parser.add_argument("--end", type=int, default=None, help="Nomor baris akhir")
    parser.add_argument("--limit", type=int, default=None, help="Jumlah baris maksimum yang diproses")
    parser.add_argument("--delay", type=float, default=1.5, help="Jeda antar baris dalam detik (default: 1.5)")
    parser.add_argument("--delay-kirim", type=float, default=KIRIM_DELAY, help="Jeda sebelum klik Kirim dalam detik (default: 10)")
    parser.add_argument("--force", action="store_true", help="Proses ulang semua baris meskipun sudah SUCCESS")
    parser.add_argument("--headless", action="store_true", help="Jalankan tanpa jendela Chrome")
    parser.add_argument("--dry-run", action="store_true", help="Hanya tampilkan daftar baris tanpa eksekusi")
    args = parser.parse_args()

    KIRIM_DELAY = args.delay_kirim

    source_path = Path(args.source)

    if args.init_sample:
        create_sample_source(source_path)
        return

    all_tasks = read_excel_source(source_path)
    progress_data = load_progress()

    completed = {int(k) for k, v in progress_data.items() if v.get("status") == "SUCCESS"}
    failed = {int(k) for k, v in progress_data.items() if v.get("status") == "FAILED"}

    filtered_tasks = []
    skipped = 0
    for t in all_tasks:
        idx = t["index"]
        if idx < args.start:
            continue
        if args.end and idx > args.end:
            continue
        if not args.force and idx in completed:
            skipped += 1
            continue
        filtered_tasks.append(t)
        if args.limit and len(filtered_tasks) >= args.limit:
            break

    console.print(
        Panel.fit(
            f"[bold cyan]⚡ FASIH-SM Batch Remark NIK Engine[/bold cyan]\n"
            f"File Sumber      : [yellow]{source_path}[/yellow] (Total {len(all_tasks)} baris)\n"
            f"SSO User         : [bold green]{config.SSO_USERNAME}[/bold green] (Auto-Login: {config.AUTO_LOGIN_SSO})\n"
            f"Profil Browser   : [magenta]{config.BROWSER_DATA_DIR}[/magenta]\n"
            f"Sudah Sukses     : [bold green]{len(completed)} baris[/bold green] (Otomatis dilewati: {skipped})\n"
            f"Pernah Gagal     : [bold red]{len(failed)} baris[/bold red] (Akan dicoba ulang)\n"
            f"Target Proses Ini: [bold cyan]{len(filtered_tasks)} baris[/bold cyan] (Mulai baris {args.start})\n"
            f"Tampilan Visual  : [yellow]{'Headless' if args.headless else 'Google Chrome Jendela Visual'}[/yellow] | Delay: [cyan]{args.delay}s[/cyan]",
            border_style="cyan",
        )
    )

    if not filtered_tasks:
        console.print(
            "[bold green]✔ Semua baris target sudah SUCCESS sebelumnya![/bold green]\n"
            "[dim]Gunakan --force jika ingin memproses ulang.[/dim]"
        )
        return

    if args.dry_run:
        table = Table(title="Daftar Baris Target (Dry Run)")
        for name, style in [("No", "cyan"), ("Nama ART", "magenta"), ("NIK", "green"),
                            ("Komentar", "yellow"), ("Timpa", "red"), ("Status Terakhir", "yellow")]:
            table.add_column(name, style=style, justify="right" if name == "No" else "left")
        for t in filtered_tasks[:20]:
            prev = progress_data.get(str(t["index"]), {}).get("status", "PENDING")
            table.add_row(str(t["index"]), t["nama"], t["nik"] or "-", t["komentar"],
                          "ya" if _is_true(t["timpa_remark"]) else "-", prev)
        console.print(table)
        if len(filtered_tasks) > 20:
            console.print(f"[dim]... dan {len(filtered_tasks) - 20} baris lainnya.[/dim]")
        return

    console.print(f"[cyan]🚀 Meluncurkan Google Chrome otomatis ({'headless' if args.headless else 'visual'})...[/cyan]")
    config.BROWSER_DATA_DIR.mkdir(parents=True, exist_ok=True)

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(config.BROWSER_DATA_DIR),
            channel=config.BROWSER_CHANNEL,
            headless=args.headless,
            viewport=None,
            args=["--start-maximized", "--disable-blink-features=AutomationControlled"],
        )
        page = context.pages[0] if context.pages else context.new_page()

        try:
            page.goto(f"{config.FASIH_BASE_URL}/app/", wait_until="domcontentloaded", timeout=25000)
            time.sleep(2.0)
            if is_page_logged_out(page):
                console.print("[yellow]⚠ Status awal logout! Login otomatis SSO BPS...[/yellow]")
                if not perform_sso_login(page):
                    console.print("[bold red]❌ Gagal login otomatis SSO. Proses dibatalkan.[/bold red]")
                    context.close()
                    return
        except Exception as e:
            console.print(f"[dim yellow]Peringatan cek awal: {e}[/dim yellow]")

        success_count = 0
        fail_count = 0
        current_item = None

        def record(item: dict, status: str, detail: str):
            progress_data[str(item["index"])] = {
                "link": item["link"], "nama": item["nama"], "nik": item["nik"],
                "komentar": item["komentar"], "status": status, "detail": detail,
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            }
            save_progress(progress_data)

        try:
            for i, item in enumerate(filtered_tasks, start=1):
                current_item = item
                idx = item["index"]

                console.print(f"\n[bold blue]━━━ [{i}/{len(filtered_tasks)}] Memproses Baris #{idx} ━━━[/bold blue]")
                console.print(f"ART     : [magenta]{item['nama']}[/magenta] | NIK: [cyan]{item['nik'] or '-'}[/cyan] | Komentar: [cyan]{item['komentar']}[/cyan]")
                console.print(f"Link    : [dim]{item['link']}[/dim]")

                try:
                    is_ok, msg = remark_nik_row(
                        page, item["link"], item["nama"], item["komentar"],
                        timpa=_is_true(item["timpa_remark"]),
                    )
                except KeyboardInterrupt:
                    raise
                except Exception as e:
                    is_ok, msg = False, f"Exception tak terduga: {e}"

                record(item, "SUCCESS" if is_ok else "FAILED", msg)
                current_item = None

                if is_ok:
                    success_count += 1
                    console.print(f"[bold green]✔ #{idx} Berhasil: {msg}[/bold green]")
                else:
                    fail_count += 1
                    console.print(f"[bold red]❌ #{idx} Gagal: {msg}[/bold red]")

                if i % 10 == 0 or i == len(filtered_tasks):
                    export_result_excel(all_tasks, progress_data, RESULT_EXCEL_FILE)

                if i < len(filtered_tasks) and args.delay > 0:
                    time.sleep(args.delay)

        except KeyboardInterrupt:
            console.print("\n[bold yellow]⚠ Eksekusi dihentikan pengguna (Ctrl+C). Checkpoint tersimpan aman![/bold yellow]")
            if current_item is not None:
                record(current_item, "FAILED", "Dihentikan pengguna (Ctrl+C) saat sedang diproses")
                console.print(f"[yellow]↳ Baris #{current_item['index']} ditandai FAILED (terhenti) & akan dicoba ulang.[/yellow]")
        finally:
            export_result_excel(all_tasks, progress_data, RESULT_EXCEL_FILE)
            context.close()
            console.print(
                Panel.fit(
                    f"[bold cyan]📊 Rekapitulasi Remark NIK[/bold cyan]\n"
                    f"Selesai Diproses : [bold]{success_count + fail_count}[/bold] baris\n"
                    f"Berhasil (Submit): [bold green]{success_count}[/bold green]\n"
                    f"Gagal            : [bold red]{fail_count}[/bold red]\n"
                    f"Log Checkpoint   : [yellow]{PROGRESS_FILE}[/yellow]\n"
                    f"Hasil Excel      : [green]{RESULT_EXCEL_FILE}[/green]",
                    border_style="cyan",
                )
            )


if __name__ == "__main__":
    main()