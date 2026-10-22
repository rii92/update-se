#!/usr/bin/env python3
"""
Batch Automation Cleaning NIK Duplikat FASIH-SM SE2026 (Grouped by Assignment ID)
Membaca daftar assignment dan nama ART dari Excel (source/cleaning_0_nik_duplikat.xlsx):
1. Mengelompokkan baris berdasarkan ASSIGNMENT ID yang sama.
2. Otomatis membuka Google Chrome dengan profil terisolasi (./chrome_data).
3. Otomatis login SSO BPS jika belum login / sesi berakhir.
4. Membuka link kuesioner (/edit) HANYA SEKALI per assignment.
5. Masuk ke section "SE2026 - L BLOK I".
6. Memberikan komentar "#0" pada NIK untuk SEMUA ART yang terdaftar di assignment tersebut.
7. Masuk ke section CATATAN, memastikan switch anomali aktif (jika belum), dan submit sekali.
8. Merekap progres per baris ke source/cleaning_nik_progress.json dan source/cleaning_nik_result.xlsx.
"""

import argparse
import json
import os
import sys
import time
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from playwright.sync_api import BrowserContext, Page, sync_playwright
from rich.console import Console
from rich.panel import Panel
from rich.table import Table

import config

console = Console()

PROGRESS_FILE = Path("source/cleaning_nik_progress.json")
RESULT_EXCEL_FILE = Path("source/cleaning_nik_result.xlsx")


def load_progress() -> Dict[str, dict]:
    """Memuat data checkpoint dari file json dan file Excel hasil (jika ada)."""
    progress = {}

    # 1. Dari JSON jika ada
    if PROGRESS_FILE.exists():
        try:
            with open(PROGRESS_FILE, "r") as f:
                progress = json.load(f)
        except Exception:
            pass

    # 2. Sinkronkan dari Excel jika ada data tambahan
    if RESULT_EXCEL_FILE.exists():
        try:
            wb = openpyxl.load_workbook(RESULT_EXCEL_FILE, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                idx_str = str(row[0]).strip()
                status = str(row[4]).strip() if len(row) > 4 and row[4] else ""
                if idx_str not in progress and status:
                    progress[idx_str] = {
                        "link": str(row[1]).strip() if len(row) > 1 and row[1] else "",
                        "nama_art": str(row[2]).strip() if len(row) > 2 and row[2] else "",
                        "komentar": str(row[3]).strip() if len(row) > 3 and row[3] else "#0",
                        "status": status,
                        "detail": str(row[5]).strip() if len(row) > 5 and row[5] else "",
                        "timestamp": str(row[6]).strip() if len(row) > 6 and row[6] else "",
                    }
        except Exception:
            pass

    return progress


def save_progress(data: Dict[str, dict]):
    """Menyimpan data checkpoint ke file json."""
    PROGRESS_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(PROGRESS_FILE, "w") as f:
        json.dump(data, f, indent=2)


def export_result_excel(tasks: List[dict], progress_data: Dict[str, dict], output_path: Path):
    """Mengekspor rekapitulasi status eksekusi per baris ke Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hasil Cleaning NIK"

    headers = ["No", "Link", "NAMA ART SE", "Komentar", "Status", "Detail", "Waktu Eksekusi"]
    ws.append(headers)

    for item in tasks:
        idx = item["index"]
        rec = progress_data.get(str(idx), {})
        status = rec.get("status", "PENDING")
        detail = rec.get("detail", "")
        ts = rec.get("timestamp", "")
        ws.append([idx, item["link"], item["nama_art"], item["komentar"], status, detail, ts])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def read_excel_source(file_path: Path) -> List[dict]:
    """Membaca daftar assignment dan nama ART dari file Excel."""
    if not file_path.exists():
        console.print(f"[bold red]❌ File sumber tidak ditemukan: {file_path}[/bold red]")
        sys.exit(1)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active

    headers = [str(c).strip() if c is not None else "" for c in next(sheet.iter_rows(values_only=True))]

    # Cari indeks kolom yang relevan
    link_col = None
    nama_col = None
    aid_col = None

    for i, h in enumerate(headers):
        h_upper = h.upper()
        if "LINK FASIH" in h_upper or h_upper == "LINK":
            link_col = i
        elif "NAMA ART SE" in h_upper:
            nama_col = i
        elif "ASSIGNMENT ID" in h_upper:
            aid_col = i

    if link_col is None:
        link_col = 6  # fallback default indeks 6
    if nama_col is None:
        nama_col = 12  # fallback default indeks 12
    if aid_col is None:
        aid_col = 3  # fallback default indeks 3

    tasks = []
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
        if not row or len(row) <= link_col or not row[link_col]:
            continue
        link = str(row[link_col]).strip()
        nama_art = str(row[nama_col]).strip() if len(row) > nama_col and row[nama_col] is not None else ""
        aid = str(row[aid_col]).strip() if len(row) > aid_col and row[aid_col] is not None else link.split("/")[-1]

        tasks.append({
            "index": row_idx,
            "assignment_id": aid,
            "link": link,
            "nama_art": nama_art,
            "komentar": "#0",
        })

    return tasks


def group_tasks_by_assignment(tasks: List[dict]) -> List[dict]:
    """
    Mengelompokkan daftar baris task berdasarkan Assignment ID.
    Menjaga urutan kemunculan pertama assignment pada file sumber.
    """
    groups = OrderedDict()
    for t in tasks:
        aid = t.get("assignment_id") or t["link"].strip().rstrip("/").removesuffix("/edit").split("/")[-1]
        if aid not in groups:
            groups[aid] = {
                "assignment_id": aid,
                "link": t["link"],
                "first_index": t["index"],
                "tasks": [],
            }
        groups[aid]["tasks"].append(t)
    return list(groups.values())


def is_page_logged_out(page: Page) -> bool:
    """Mengecek apakah browser sedang berada pada status logout atau SSO."""
    try:
        url = page.url.lower()
        if "sso.bps.go.id" in url or "/login" in url or "/auth/login" in url or "/oauth2/authorization" in url:
            return True
        if page.locator("#kc-form-login, #username, #password, .login-button").count() > 0:
            return True
        text = page.locator("body").inner_text() if page.locator("body").count() > 0 else ""
        if "Selamat Datang Kembali" in text or "Lanjutkan dengan SSO" in text or "Login SSO BPS" in text:
            return True
        return False
    except Exception:
        return False


def perform_sso_login(page: Page, username: str = config.SSO_USERNAME, password: str = config.SSO_PASSWORD) -> bool:
    """Melakukan login otomatis SSO BPS."""
    console.print(f"[bold yellow]🔐 Login otomatis SSO BPS (User: {username})...[/bold yellow]")
    try:
        if "sso.bps.go.id" not in page.url:
            sso_btn = page.locator(
                "a:has-text('Lanjutkan dengan SSO'), a:has-text('Login SSO BPS'), a[href*='oauth2/authorization'], .login-button"
            )
            if sso_btn.count() > 0 and sso_btn.first.is_visible():
                sso_btn.first.click()
            else:
                page.goto("https://fasih-sm.bps.go.id/oauth2/authorization/ics", wait_until="domcontentloaded")
            page.wait_for_url("**/sso.bps.go.id/**", timeout=20000)

        page.wait_for_selector("#username", timeout=15000)
        page.fill("#username", username)
        page.fill("#password", password)
        page.click("#kc-login")

        page.wait_for_url("**/fasih-sm.bps.go.id/app/**", timeout=25000)
        console.print(f"[bold green]✔ Berhasil login SSO BPS sebagai {username}![/bold green]")
        time.sleep(2.0)
        return True
    except Exception as e:
        console.print(f"[bold red]❌ Gagal login otomatis SSO: {e}[/bold red]")
        return False


def comment_art_nik(page: Page, nama_art: str, komentar_value: str = "#0") -> Tuple[bool, str]:
    """
    Mencari ART di sub-section SE2026 - L BLOK I, membuka dialog catatan NIK, dan mengisi komentar.
    """
    target_name_clean = nama_art.strip().lower()

    # 1. Cari nama ART di aside dan klik
    click_art_res = page.evaluate(f"""(() => {{
        const aside = document.querySelector("aside");
        if (!aside) return "ERR_NO_ASIDE";
        
        const clean = (s) => (s || "").toLowerCase().replace(/\\s+/g, ' ').trim();
        const target = "{target_name_clean}";
        
        let items = Array.from(aside.querySelectorAll("*")).filter(el => 
            el.children.length === 0 && el.textContent.trim().length > 0
        );

        let match = items.find(el => clean(el.textContent) === target);
        if (!match) {{
            match = items.find(el => clean(el.textContent).includes(target) || (target.length > 4 && clean(el.textContent).length > 4 && target.includes(clean(el.textContent))));
        }}

        // Jika belum ketemu, coba re-expand SE2026 - L BLOK I
        if (!match) {{
            const blok = items.find(el => clean(el.textContent).includes("se2026 - l blok i"));
            if (blok) {{
                let clickableBlok = blok;
                while (clickableBlok && clickableBlok !== aside && !clickableBlok.onclick && !clickableBlok.className.includes("cursor-pointer")) {{
                    clickableBlok = clickableBlok.parentElement;
                }}
                (clickableBlok || blok).click();
                return "RE_EXPANDED";
            }}
        }}

        if (!match) return "NOT_FOUND";

        let clickable = match;
        while (clickable && clickable !== aside && !clickable.onclick && !clickable.className.includes("cursor-pointer")) {{
            clickable = clickable.parentElement;
        }}
        (clickable || match).scrollIntoView({{ behavior: "instant", block: "center" }});
        (clickable || match).click();
        return "CLICKED";
    }})()""")

    # Jika sempat re-expand, coba cari sekali lagi
    if click_art_res == "RE_EXPANDED":
        time.sleep(1.5)
        click_art_res = page.evaluate(f"""(() => {{
            const aside = document.querySelector("aside");
            if (!aside) return "ERR_NO_ASIDE";
            const clean = (s) => (s || "").toLowerCase().replace(/\\s+/g, ' ').trim();
            const target = "{target_name_clean}";
            const items = Array.from(aside.querySelectorAll("*")).filter(el => 
                el.children.length === 0 && el.textContent.trim().length > 0
            );
            let match = items.find(el => clean(el.textContent) === target);
            if (!match) {{
                match = items.find(el => clean(el.textContent).includes(target) || (target.length > 4 && clean(el.textContent).length > 4 && target.includes(clean(el.textContent))));
            }}
            if (!match) return "NOT_FOUND";
            let clickable = match;
            while (clickable && clickable !== aside && !clickable.onclick && !clickable.className.includes("cursor-pointer")) {{
                clickable = clickable.parentElement;
            }}
            (clickable || match).scrollIntoView({{ behavior: "instant", block: "center" }});
            (clickable || match).click();
            return "CLICKED";
        }})()""")

    if click_art_res != "CLICKED":
        return False, f"Nama ART '{nama_art}' tidak ditemukan di sub-section SE2026 - L BLOK I"

    time.sleep(2.0)

    # 2. Cari rincian pertanyaan "Nomor Induk Kependudukan (NIK)"
    find_nik_script = """(() => {
        const findNikRow = () => {
            const candidateDivs = Array.from(document.querySelectorAll("div[class*='rounded-lg'], div[class*='flex-row'], div[class*='gap-2']"));
            let row = candidateDivs.find(r => {
                const t = (r.innerText || "").toLowerCase();
                return (t.includes("nomor induk kependudukan") || t.includes("nik")) && !t.includes("cek nik") && (r.querySelector("button[title='Catatan'], button[aria-label='Catatan']") !== null);
            });
            if (!row) {
                const allDivs = Array.from(document.querySelectorAll("div"));
                row = allDivs.find(r => {
                    const t = (r.innerText || "").toLowerCase();
                    return (t.includes("nomor induk kependudukan") || t.includes("nik")) && !t.includes("cek nik") && (r.querySelector("button[title='Catatan'], button[aria-label='Catatan']") !== null);
                });
            }
            return row;
        };

        const nikRow = findNikRow();
        if (!nikRow) return "ERR_NIK_ROW_NOT_FOUND";

        const btns = Array.from(nikRow.querySelectorAll("button[title='Catatan'], button[aria-label='Catatan']"));
        const visibleBtn = btns.find(b => b.offsetWidth > 0 && b.offsetHeight > 0) || btns[0];
        if (!visibleBtn) return "ERR_NO_CATATAN_BTN";

        visibleBtn.scrollIntoView({ behavior: "smooth", block: "center" });
        visibleBtn.click();
        return "CLICKED";
    })()"""

    click_catatan_res = page.evaluate(find_nik_script)

    if click_catatan_res == "ERR_NIK_ROW_NOT_FOUND":
        time.sleep(1.0)
        # Coba sekali lagi barangkali form baru saja selesai render
        click_catatan_res = page.evaluate(find_nik_script)

    if click_catatan_res != "CLICKED":
        return False, f"Gagal mengklik tombol Catatan pada rincian NIK ({click_catatan_res})"

    try:
        page.wait_for_selector("textarea[placeholder*='catatan' i]", timeout=8000)
    except Exception:
        return False, "Dialog catatan NIK tidak muncul"

    # 3. Isi komentar dan simpan
    save_res = page.evaluate(f"""(() => {{
        const ta = document.querySelector("textarea[placeholder*='catatan' i]");
        if (!ta) return "ERR_NO_TEXTAREA";

        const dialog = ta.closest("[role='dialog']") || document.querySelector("[class*='max-w-lg']");
        const valToFill = "{komentar_value}";

        // Jika komentar sudah ada
        if (dialog && dialog.innerText.includes(valToFill)) {{
            const closeBtn = dialog.querySelector("button[aria-label='Dismiss']") || 
                             Array.from(dialog.querySelectorAll("button")).find(b => b.innerText.trim() === "Close" || b.innerText.trim() === "Batal");
            if (closeBtn) closeBtn.click();
            return "ALREADY_EXISTS";
        }}

        ta.focus();
        const nativeSetter = Object.getOwnPropertyDescriptor(window.HTMLTextAreaElement.prototype, "value").set;
        nativeSetter.call(ta, valToFill);
        ta.dispatchEvent(new Event("input", {{ bubbles: true }}));
        ta.dispatchEvent(new Event("change", {{ bubbles: true }}));

        const saveBtn = document.querySelector("button[title='Simpan']");
        if (!saveBtn) return "ERR_NO_SAVE_BTN";
        saveBtn.click();

        setTimeout(() => {{
            const closeBtn = document.querySelector("button[aria-label='Dismiss']") || 
                             Array.from(document.querySelectorAll("button")).find(b => b.innerText.trim() === "Close" || b.innerText.trim() === "Batal");
            if (closeBtn) closeBtn.click();
        }}, 600);

        return "SAVED";
    }})()""")

    # Pastikan modal dialog tertutup rapat agar tidak menghalangi klik berikutnya
    try:
        page.wait_for_selector("[role='dialog']", state="hidden", timeout=3000)
    except Exception:
        page.evaluate("""(() => {
            const d = document.querySelector("[role='dialog']");
            if (d) {
                const b = d.querySelector("button[aria-label='Dismiss']") || 
                          Array.from(d.querySelectorAll("button")).find(x => x.innerText.trim() === 'Close' || x.innerText.trim() === 'Batal');
                if (b) b.click();
            }
        })()""")
        time.sleep(0.5)

    if save_res in ("SAVED", "ALREADY_EXISTS"):
        return True, f"Komentar {komentar_value} {'sudah ada' if save_res == 'ALREADY_EXISTS' else 'disimpan'}"
    else:
        return False, f"Gagal menyimpan catatan: {save_res}"


def submit_assignment(page: Page) -> Tuple[bool, str]:
    """
    Melakukan proses submit kuesioner:
    1. Masuk ke section CATATAN
    2. Cek switch "Tampilkan Anomali Usaha dan Keluarga" (aktifkan jika belum)
    3. Pastikan switch "Anomali diselesaikan oleh admin" tidak aktif
    4. Klik Kirim utama -> Cek validasi 0 Galat -> Klik Kirim modal -> Konfirmasi final
    """
    # 1. Masuk ke section CATATAN
    click_catatan_ok = page.evaluate("""(() => {
        const aside = document.querySelector("aside");
        if (!aside) return false;
        const elements = Array.from(aside.querySelectorAll("*")).filter(el => el.children.length === 0 && el.textContent.trim() === "CATATAN");
        if (elements.length === 0) return false;
        const target = elements[0];
        let clickable = target;
        while (clickable && clickable !== aside && !clickable.onclick && !clickable.className.includes("cursor-pointer")) {
            clickable = clickable.parentElement;
        }
        clickable.scrollIntoView({ block: "center" });
        (clickable || target).click();
        return true;
    })()""")

    if not click_catatan_ok:
        return False, "Menu CATATAN tidak ditemukan di sidebar"

    time.sleep(1.5)

    # 2. Cek switch di blok CATATAN:
    # - "Tampilkan Anomali Usaha dan Keluarga" jika sudah dichecklist JANGAN dichecklist lagi.
    # - JANGAN checklist "Anomali diselesaikan oleh admin".
    # - JANGAN memberi catatan penjelasan anomali admin.
    page.evaluate("""(() => {
        const rows = Array.from(document.querySelectorAll("div[class*='rounded-lg'], div[class*='gap-2'], div[class*='flex'], label"));

        // Cari switch "Tampilkan Anomali Usaha dan Keluarga" (hindari baris admin)
        const anomaliRow = rows.find(r => {
            const t = (r.innerText || "").toLowerCase();
            return t.includes("tampilkan anomali") && !t.includes("admin");
        });

        if (anomaliRow) {
            const input = anomaliRow.querySelector("input[role='switch'], input[type='checkbox']");
            if (input) {
                const isChecked = input.checked || 
                                  input.getAttribute("aria-checked") === "true" || 
                                  input.hasAttribute("data-checked");
                if (!isChecked) {
                    const control = anomaliRow.querySelector("[id*='control']") || input;
                    control.click();
                }
            }
        }

        // Pastikan switch "Anomali diselesaikan oleh admin" TIDAK DICHECKLIST (jika sempat aktif, uncheck)
        const adminRow = rows.find(r => {
            const t = (r.innerText || "").toLowerCase();
            return t.includes("anomali diselesaikan oleh admin");
        });
        if (adminRow) {
            const adminInput = adminRow.querySelector("input[role='switch'], input[type='checkbox']");
            if (adminInput) {
                const isAdminChecked = adminInput.checked || 
                                       adminInput.getAttribute("aria-checked") === "true" || 
                                       adminInput.hasAttribute("data-checked");
                if (isAdminChecked) {
                    const adminControl = adminRow.querySelector("[id*='control']") || adminInput;
                    adminControl.click();
                }
            }
        }
    })()""")

    time.sleep(1.0)

    # 3. Klik tombol Kirim utama
    page.evaluate("""(() => {
        const btn = Array.from(document.querySelectorAll("button")).find(b => b.innerText.trim() === "Kirim" && b.offsetWidth > 0);
        if (btn) {
            btn.scrollIntoView({ block: "center" });
            btn.click();
        }
    })()""")

    try:
        page.wait_for_selector("[class*='max-w-lg'], [role='dialog']", timeout=10000)
    except Exception:
        return False, "Modal ringkasan validasi tidak muncul"

    time.sleep(1.0)

    # 4. Validasi 0 Galat & Klik Kirim di modal
    val_res = page.evaluate("""(() => {
        const modal = document.querySelector("[class*='max-w-lg'], [role='dialog']");
        if (!modal) return { status: "ERR", msg: "Modal validasi tidak ditemukan" };
        const modalText = modal.innerText;
        const galatMatch = modalText.match(/GALAT\\s*\\n\\s*(\\d+)/i);
        const galatCount = galatMatch ? parseInt(galatMatch[1]) : 0;
        if (galatCount > 0) {
            const closeBtn = modal.querySelector("button[aria-label='Dismiss']") || 
                             Array.from(modal.querySelectorAll("button")).find(b => b.innerText.trim() === "Batal");
            if (closeBtn) closeBtn.click();
            return { status: "GALAT", count: galatCount, msg: galatCount + " galat ditemukan!" };
        }
        const btns = Array.from(modal.querySelectorAll("button")).filter(b => b.innerText.trim() === "Kirim");
        if (btns.length === 0) return { status: "ERR", msg: "Tombol Kirim di modal validasi tidak ditemukan" };
        btns[0].click();
        return { status: "OK" };
    })()""")

    if val_res.get("status") != "OK":
        return False, f"Validasi gagal: {val_res.get('msg')}"

    # 5. Klik tombol Konfirmasi pada dialog akhir
    try:
        page.wait_for_selector("button:has-text('Konfirmasi')", timeout=10000)
        page.locator("button:has-text('Konfirmasi')").first.click()
    except Exception:
        return False, "Tombol 'Konfirmasi' dialog akhir tidak ditemukan"

    # 6. Verifikasi submit selesai
    try:
        page.wait_for_function("""() => {
            const t = document.body ? document.body.innerText : "";
            return t.includes("Assignment Berhasil Disubmit") || t.includes("Perubahan Anda berhasil disimpan");
        }""", timeout=15000)
        return True, "Assignment Berhasil Disubmit"
    except Exception:
        return False, "Status sukses submit tidak terdeteksi pada layar"


def process_assignment_nik(
    page: Page,
    assignment: dict,
    pending_tasks: List[dict],
    retry_on_logout: bool = True,
) -> Dict[int, Tuple[bool, str]]:
    """
    Membuka kuesioner satu kali, memproses semua pending ART di dalamnya, lalu submit.
    Mengembalikan mapping { task_index: (is_success, detail_message) }.
    """
    link = assignment["link"]
    edit_url = link.strip()
    if not edit_url.endswith("/edit"):
        edit_url = edit_url.rstrip("/") + "/edit"

    try:
        page.goto(edit_url, wait_until="domcontentloaded", timeout=30000)
    except Exception as e:
        return {t["index"]: (False, f"Timeout membuka URL: {e}") for t in pending_tasks}

    time.sleep(1.0)

    # Cek jika sesi logout
    if is_page_logged_out(page):
        if retry_on_logout and perform_sso_login(page):
            return process_assignment_nik(page, assignment, pending_tasks, retry_on_logout=False)
        return {t["index"]: (False, "Sesi login habis / terlogout") for t in pending_tasks}

    # Tunggu aside menu kuesioner siap
    try:
        page.wait_for_selector("aside", timeout=20000)
    except Exception:
        if is_page_logged_out(page) and retry_on_logout:
            if perform_sso_login(page):
                return process_assignment_nik(page, assignment, pending_tasks, retry_on_logout=False)
        return {t["index"]: (False, "Menu sidebar (aside) kuesioner tidak muncul") for t in pending_tasks}

    time.sleep(1.0)

    # Masuk ke blok "SE2026 - L BLOK I"
    click_blok_ok = page.evaluate("""(() => {
        const aside = document.querySelector("aside");
        if (!aside) return false;
        const items = Array.from(aside.querySelectorAll("*")).filter(el => 
            el.children.length === 0 && (el.textContent.trim() === "SE2026 - L BLOK I" || el.textContent.includes("SE2026 - L BLOK I"))
        );
        if (items.length === 0) return false;
        const target = items[0];
        let clickable = target;
        while (clickable && clickable !== aside && !clickable.onclick && !clickable.className.includes("cursor-pointer")) {
            clickable = clickable.parentElement;
        }
        clickable.scrollIntoView({ block: 'center' });
        (clickable || target).click();
        return true;
    })()""")

    if not click_blok_ok:
        return {t["index"]: (False, "Blok SE2026 - L BLOK I tidak ditemukan di sidebar") for t in pending_tasks}

    time.sleep(2.0)

    # Loop setiap pending ART dalam assignment ini
    art_results: Dict[int, Tuple[bool, str]] = {}
    name_cache: Dict[str, Tuple[bool, str]] = {}

    for t in pending_tasks:
        idx = t["index"]
        nama_art = t["nama_art"]
        clean_name = nama_art.strip().lower()

        # Jika nama ART yang sama sudah diproses dalam assignment ini (misal ada duplikasi baris dengan ART sama)
        if clean_name in name_cache:
            art_results[idx] = name_cache[clean_name]
            console.print(f"   [dim]↳ Baris #{idx} ({nama_art}): Menggunakan hasil komentar sebelumnya[/dim]")
            continue

        console.print(f"   [cyan]👉 Baris #{idx} - Mengomentari ART:[/cyan] [bold]{nama_art}[/bold]...")
        is_ok, msg = comment_art_nik(page, nama_art, t.get("komentar", "#0"))
        name_cache[clean_name] = (is_ok, msg)
        art_results[idx] = (is_ok, msg)

        if is_ok:
            console.print(f"   [green]✔ Baris #{idx} ({nama_art}): {msg}[/green]")
        else:
            console.print(f"   [yellow]⚠ Baris #{idx} ({nama_art}): {msg}[/yellow]")

        time.sleep(1.0)

    # Cek apakah ada minimal 1 ART yang berhasil dikomentari
    any_success = any(res[0] for res in art_results.values())
    if not any_success:
        console.print("   [bold red]❌ Tidak ada ART yang berhasil dikomentari, lewati proses submit.[/bold red]")
        return art_results

    # Submit kuesioner satu kali untuk seluruh ART dalam assignment ini
    console.print("   [bold cyan]📤 Melakukan Submit Assignment (CATATAN -> Kirim -> Konfirmasi)...[/bold cyan]")
    submit_ok, submit_msg = submit_assignment(page)

    final_results: Dict[int, Tuple[bool, str]] = {}
    if submit_ok:
        console.print(f"   [bold green]✔ Sukses Submit Assignment: {submit_msg}[/bold green]")
        for t in pending_tasks:
            idx = t["index"]
            comment_ok, comment_detail = art_results.get(idx, (False, "Belum diproses"))
            if comment_ok:
                final_results[idx] = (True, f"Komentar {t.get('komentar', '#0')} disimpan & Assignment Berhasil Disubmit ({t['nama_art']})")
            else:
                final_results[idx] = (False, comment_detail)
    else:
        console.print(f"   [bold red]❌ Gagal Submit Assignment: {submit_msg}[/bold red]")
        for t in pending_tasks:
            idx = t["index"]
            final_results[idx] = (False, f"Gagal submit assignment: {submit_msg}")

    return final_results


def main():
    parser = argparse.ArgumentParser(description="Batch Automation Cleaning NIK Duplikat FASIH-SM SE2026 (Grouped by Assignment ID)")
    parser.add_argument("--source", default="source/cleaning_0_nik_duplikat.xlsx", help="File Excel sumber NIK duplikat")
    parser.add_argument("--start", type=int, default=1, help="Filter: Nomor baris awal untuk mulai diproses (default: 1)")
    parser.add_argument("--end", type=int, default=None, help="Filter: Nomor baris akhir untuk diproses")
    parser.add_argument("--limit", type=int, default=None, help="Batas maksimum jumlah assignment yang diproses")
    parser.add_argument("--delay", type=float, default=1.5, help="Jeda waktu antar assignment dalam detik (default: 1.5)")
    parser.add_argument("--force", action="store_true", help="Paksa proses ulang assignment meskipun sudah SUCCESS sebelumnya")
    parser.add_argument("--headless", action="store_true", help="Jalankan di latar belakang tanpa membuka jendela visual Chrome")
    parser.add_argument("--dry-run", action="store_true", help="Hanya tampilkan daftar assignment & ART tanpa membuka browser")
    args = parser.parse_args()

    source_path = Path(args.source)
    all_tasks = read_excel_source(source_path)
    total_tasks = len(all_tasks)

    # Kelompokkan berdasarkan Assignment ID
    all_assignments = group_tasks_by_assignment(all_tasks)
    total_assignments = len(all_assignments)

    progress_data = load_progress()

    completed_indices = {int(k) for k, v in progress_data.items() if v.get("status") == "SUCCESS"}
    failed_indices = {int(k) for k, v in progress_data.items() if v.get("status") == "FAILED"}

    # Saring assignment yang memiliki pending task
    candidate_assignments = []
    skipped_assignments = 0
    total_pending_rows = 0

    for grp in all_assignments:
        # Cek task mana saja yang belum SUCCESS
        pending = [
            t for t in grp["tasks"]
            if args.force or progress_data.get(str(t["index"]), {}).get("status") != "SUCCESS"
        ]

        if not pending:
            skipped_assignments += 1
            continue

        # Terapkan filter start/end jika ada (berdasarkan nomor baris task)
        if args.start > 1:
            pending = [t for t in pending if t["index"] >= args.start]
        if args.end:
            pending = [t for t in pending if t["index"] <= args.end]

        if not pending:
            continue

        candidate_assignments.append({
            "assignment": grp,
            "pending_tasks": pending,
        })
        total_pending_rows += len(pending)

        if args.limit and len(candidate_assignments) >= args.limit:
            break

    console.print(
        Panel.fit(
            f"[bold cyan]⚡ FASIH-SM Auto-Cleaning NIK Duplikat (Grouped by Assignment ID)[/bold cyan]\n"
            f"File Sumber          : [yellow]{source_path}[/yellow] ({total_tasks} baris | {total_assignments} unique assignment)\n"
            f"SSO User             : [bold green]{config.SSO_USERNAME}[/bold green] (Auto-Login: {config.AUTO_LOGIN_SSO})\n"
            f"Profil Browser       : [magenta]{config.BROWSER_DATA_DIR}[/magenta]\n"
            f"Baris Sudah Sukses   : [bold green]{len(completed_indices)} baris[/bold green] (Assignment dilewati: {skipped_assignments})\n"
            f"Baris Pernah Gagal   : [bold red]{len(failed_indices)} baris[/bold red] (Akan dicoba ulang)\n"
            f"Target Eksekusi Ini  : [bold cyan]{len(candidate_assignments)} assignment[/bold cyan] (Mencakup [bold yellow]{total_pending_rows} baris ART[/bold yellow])\n"
            f"Tampilan Visual      : [yellow]{'Headless (Latar Belakang)' if args.headless else 'Google Chrome Jendela Visual'}[/yellow] | Delay: [cyan]{args.delay}s[/cyan]",
            border_style="cyan",
        )
    )

    if not candidate_assignments:
        console.print(
            "[bold green]✔ Semua assignment target sudah SUCCESS sebelumnya! Tidak ada yang perlu diproses.[/bold green]\n"
            "[dim]Gunakan opsi --force jika ingin tetap memproses ulang yang sudah sukses.[/dim]"
        )
        return

    if args.dry_run:
        table = Table(title="Daftar Target Assignment & ART (Dry Run)")
        table.add_column("No", justify="right", style="cyan")
        table.add_column("Assignment ID", style="magenta")
        table.add_column("Jml ART", justify="center", style="yellow")
        table.add_column("Daftar ART (No Baris: Nama)", style="bold green")

        for i, item in enumerate(candidate_assignments[:25], start=1):
            grp = item["assignment"]
            pending = item["pending_tasks"]
            art_info = ", ".join([f"#{t['index']} {t['nama_art']}" for t in pending])
            if len(art_info) > 60:
                art_info = art_info[:57] + "..."
            table.add_row(str(i), grp["assignment_id"], str(len(pending)), art_info)

        console.print(table)
        if len(candidate_assignments) > 25:
            console.print(f"[dim]... dan {len(candidate_assignments) - 25} assignment lainnya.[/dim]")
        return

    console.print(f"[cyan]🚀 Meluncurkan Google Chrome otomatis ({'visual' if not args.headless else 'headless'})...[/cyan]")
    config.BROWSER_DATA_DIR.mkdir(parents=True, exist_ok=True)

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(config.BROWSER_DATA_DIR),
            channel=config.BROWSER_CHANNEL,
            headless=args.headless,
            viewport=None,
            args=[
                "--start-maximized",
                "--disable-blink-features=AutomationControlled",
            ],
        )

        page = context.pages[0] if context.pages else context.new_page()

        # Pastikan status login awal di dashboard
        try:
            page.goto("https://fasih-sm.bps.go.id/app/", wait_until="domcontentloaded", timeout=25000)
            time.sleep(2.0)
            if is_page_logged_out(page):
                console.print("[yellow]⚠ Status awal terdeteksi logout! Melakukan login otomatis SSO BPS...[/yellow]")
                if not perform_sso_login(page):
                    console.print("[bold red]❌ Gagal login otomatis SSO. Proses dibatalkan.[/bold red]")
                    context.close()
                    return
        except Exception as e:
            console.print(f"[dim yellow]Peringatan cek awal: {e}[/dim yellow]")

        total_processed_assignments = 0
        total_success_rows = 0
        total_failed_rows = 0

        try:
            for i, cand in enumerate(candidate_assignments, start=1):
                grp = cand["assignment"]
                pending = cand["pending_tasks"]
                aid = grp["assignment_id"]
                link = grp["link"]

                art_summary = ", ".join([f"#{t['index']} {t['nama_art']}" for t in pending])
                console.print(f"\n[bold blue]━━━ [Assignment {i}/{len(candidate_assignments)}] ID: {aid} ({len(pending)} ART) ━━━[/bold blue]")
                console.print(f"[dim]Link: {link}[/dim]")
                console.print(f"Target ART: [bold green]{art_summary}[/bold green]")

                # Jalankan proses 1 assignment untuk semua pending ART
                results = process_assignment_nik(page, grp, pending)

                curr_time_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

                # Update progress_data untuk setiap baris
                for t in pending:
                    idx = t["index"]
                    idx_str = str(idx)
                    is_ok, msg = results.get(idx, (False, "Tidak diproses"))
                    status_str = "SUCCESS" if is_ok else "FAILED"

                    progress_data[idx_str] = {
                        "link": t["link"],
                        "nama_art": t["nama_art"],
                        "komentar": t.get("komentar", "#0"),
                        "status": status_str,
                        "detail": msg,
                        "timestamp": curr_time_str,
                    }

                    if is_ok:
                        total_success_rows += 1
                    else:
                        total_failed_rows += 1

                save_progress(progress_data)
                total_processed_assignments += 1

                # Simpan berkala ke Excel setiap 5 assignment atau di akhir
                if total_processed_assignments % 5 == 0 or total_processed_assignments == len(candidate_assignments):
                    export_result_excel(all_tasks, progress_data, RESULT_EXCEL_FILE)

                if i < len(candidate_assignments) and args.delay > 0:
                    time.sleep(args.delay)

        except KeyboardInterrupt:
            console.print("\n[bold yellow]⚠ Eksekusi dihentikan pengguna (Ctrl+C). Checkpoint tersimpan aman![/bold yellow]")
        finally:
            export_result_excel(all_tasks, progress_data, RESULT_EXCEL_FILE)
            context.close()
            console.print(
                Panel.fit(
                    f"[bold cyan]📊 Rekapitulasi Eksekusi NIK Duplikat[/bold cyan]\n"
                    f"Assignment Selesai : [bold]{total_processed_assignments}[/bold] assignment\n"
                    f"Baris ART Berhasil : [bold green]{total_success_rows}[/bold green]\n"
                    f"Baris ART Gagal    : [bold red]{total_failed_rows}[/bold red]\n"
                    f"Log Checkpoint     : [yellow]{PROGRESS_FILE}[/yellow]\n"
                    f"Hasil Excel        : [green]{RESULT_EXCEL_FILE}[/green]",
                    border_style="cyan",
                )
            )


if __name__ == "__main__":
    main()
