#!/usr/bin/env python3
"""
Batch Automation Update KBLI FASIH-SM SE2026
Mengubah field KBLI pada section "KETERANGAN USAHA/PERUSAHAAN" berdasarkan
file Excel (source/update_kbli.xlsx) hanya dengan 1 perintah:

1. Otomatis membuka Google Chrome dengan profil terisolasi (config.BROWSER_DATA_DIR).
2. Otomatis login SSO BPS jika belum login / sesi habis.
3. Otomatis melewati baris yang sudah SUCCESS sebelumnya (checkpoint).
4. Untuk setiap baris:
   a. Membuka link kuesioner dalam mode edit (.../edit).
   b. Masuk section "KETERANGAN USAHA/PERUSAHAAN".
   c. Memilih sub-blok usaha sesuai "nama_usaha" (boleh sebagian, mis. "PEDAGANG SAYUR KELILING").
   d. Klik "DAPATKAN REKOMENDASI KBLI" -> "Pilih dari Master KBLI".
   e. Pada field KBLI (wrapper prefix: kbli#): buka dropdown, ketik kode KBLI,
      lalu pilih opsi "[..][kode]...".
   f. Catatan pada field KBLI (default "#1"): Catatan -> isi -> Simpan -> Dismiss.
   g. Submit assignment (masuk CATATAN -> switch anomali -> Kirim -> Konfirmasi).
5. Otomatis merekap hasil ke source/update_kbli_result.xlsx.

Kolom file sumber (source/update_kbli.xlsx):
    link          : link assignment FASIH (boleh tanpa /edit)
    nama_usaha    : nama usaha di daftar KETERANGAN USAHA/PERUSAHAAN
                    (pencocokan: persis -> diawali -> mengandung teks ini)
    kode_kbli     : kode KBLI 5 digit (mis. 47213)
    komentar      : (opsional) catatan pada field KBLI, default "#1" (nilai KBLI berubah)
    timpa_remark  : (opsional) default KBLI = catatan lama DITIMPA; isi 0 untuk mempertahankan catatan lama

Contoh:
    python update_kbli.py --init-sample
    python update_kbli.py
    python update_kbli.py --headless
    python update_kbli.py --dry-run --limit 10
    python update_kbli.py --start 1 --end 50
    python update_kbli.py --force
    python update_kbli.py --delay-kirim 15
    python update_kbli.py --no-submit --limit 1      # uji alur tanpa submit
"""

import argparse
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from playwright.sync_api import Page, sync_playwright
from rich.console import Console
from rich.panel import Panel
from rich.table import Table

import config

console = Console()

PROGRESS_FILE = Path("source/update_kbli_progress.json")
RESULT_EXCEL_FILE = Path("source/update_kbli_result.xlsx")
DEFAULT_SOURCE_FILE = Path("source/update_kbli.xlsx")

SECTION_KEYWORD = "keterangan usaha/perusahaan"
KBLI_WRAPPER_PREFIX = "kbli#"               # id wrapper field KBLI (suffix dinamis, mis. kbli#1001)
DEFAULT_KOMENTAR = "#1"                    # catatan KBLI default (nilai KBLI berubah)
KIRIM_DELAY = 10.0                         # jeda (detik) sebelum klik Kirim; ubah via --delay-kirim

_SCROLL_STEP_JS = """(step) => {
    window.scrollBy({ top: step, behavior: "instant" });
    const scrollers = Array.from(document.querySelectorAll("div")).filter(el => {
        const oy = getComputedStyle(el).overflowY;
        return (oy === "auto" || oy === "scroll") && el.scrollHeight > el.clientHeight + 20;
    });
    scrollers.forEach(el => el.scrollBy({ top: step, behavior: "instant" }));
}"""

_HAS_MORE_SCROLL_JS = """(() => {
    const scrollers = Array.from(document.querySelectorAll("div")).filter(el => {
        const oy = getComputedStyle(el).overflowY;
        return (oy === "auto" || oy === "scroll") && el.scrollHeight > el.clientHeight + 20;
    });
    return scrollers.some(el => el.scrollTop + el.clientHeight < el.scrollHeight - 5);
})()"""

_CLOSE_DIALOG_JS = """(() => {
    const d = document.querySelector("[role='dialog']") || document.querySelector("[class*='max-w-lg']");
    const scope = d || document;
    const b = scope.querySelector("button[aria-label='Dismiss']") ||
              Array.from(scope.querySelectorAll("button")).find(x => {
                  const t = x.innerText.trim();
                  return t === "Close" || t === "Batal";
              });
    if (b) { b.click(); return true; }
    return false;
})()"""


# ─────────────────────────────────────────────────────────────────────────────
# Utilitas
# ─────────────────────────────────────────────────────────────────────────────

def _is_true(val) -> bool:
    if val is None:
        return False
    return str(val).strip().lower() in ("1", "1.0", "true", "ya", "y", "yes")


def _timpa_kbli(val) -> bool:
    """KBLI: catatan lama DITIMPA secara default; kolom timpa_remark = 0 untuk mempertahankan."""
    if val is None or str(val).strip() == "":
        return True
    return _is_true(val)


def _fmt_cell(val) -> str:
    """Rapikan nilai sel Excel. Float bulat -> int; angka panjang (NIK) tidak jadi 3.2e15."""
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


# ─────────────────────────────────────────────────────────────────────────────
# Excel sumber / progress / hasil
# ─────────────────────────────────────────────────────────────────────────────

def _fmt_kbli(val: str) -> str:
    """Kode KBLI 5 digit: angka tanpa nol di depan (mis. 1111 dari Excel) dilengkapi jadi 01111."""
    s = (val or "").strip()
    if s.isdigit() and len(s) < 5:
        s = s.zfill(5)
    return s


def create_sample_source(file_path: Path = DEFAULT_SOURCE_FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Update KBLI"
    ws.append(["link", "nama_usaha", "kode_kbli", "komentar", "timpa_remark"])
    ws.append([config.DEFAULT_ASSIGNMENT_URL, "PEDAGANG SAYUR KELILING", "47213", "#1", ""])
    file_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(file_path)
    console.print(f"[bold green]✔ Contoh file sumber dibuat:[/bold green] [yellow]{file_path}[/yellow]")
    console.print(
        "[dim]Kolom: link, nama_usaha, kode_kbli, komentar (opsional), timpa_remark (opsional). "
        "Satu baris per usaha.[/dim]"
    )


def read_excel_source(file_path: Path) -> List[dict]:
    if not file_path.exists():
        console.print(f"[bold red]❌ File sumber tidak ditemukan: {file_path}[/bold red]")
        console.print("[dim]Jalankan dulu: python update_kbli.py --init-sample[/dim]")
        sys.exit(1)

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        console.print(f"[bold red]❌ File Excel tidak bisa dibaca: {file_path}[/bold red]")
        console.print(f"[dim]{e}. Pastikan formatnya .xlsx asli (bukan CSV yang di-rename) dan tidak sedang dibuka.[/dim]")
        sys.exit(1)

    sheet = wb.active
    headers = [
        str(c).strip().lower() if c is not None else ""
        for c in next(sheet.iter_rows(values_only=True))
    ]

    def exact_col(*names):
        for n in names:
            if n in headers:
                return headers.index(n)
        return None

    def contains_col(word):
        for i, h in enumerate(headers):
            if h.startswith(("remark", "timpa", "komentar", "catatan")):
                continue
            if word in h:
                return i
        return None

    link_col = exact_col("link")
    if link_col is None:
        link_col = contains_col("link")
    nama_col = exact_col("nama_usaha", "usaha", "nama")
    if nama_col is None:
        nama_col = contains_col("usaha")
    kbli_col = exact_col("kode_kbli", "kbli", "kode")
    if kbli_col is None:
        kbli_col = contains_col("kbli")
    komentar_col = exact_col("komentar", "remark", "remark_kbli", "catatan")
    timpa_col = exact_col("timpa_remark", "timpa")

    missing = [n for n, c in (("link", link_col), ("nama_usaha", nama_col), ("kode_kbli", kbli_col)) if c is None]
    if missing:
        console.print(f"[bold red]❌ Kolom wajib tidak ditemukan di header Excel: {', '.join(missing)}[/bold red]")
        sys.exit(1)

    def cell(row, col, default=""):
        if col is None or len(row) <= col or row[col] is None:
            return default
        return _fmt_cell(row[col])

    tasks = []
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
        if not row or len(row) <= link_col or not row[link_col]:
            continue
        nama = cell(row, nama_col)
        kode = _fmt_kbli(cell(row, kbli_col))
        if not nama or not kode:
            console.print(f"[dim yellow]↳ Baris #{row_idx} dilewati (nama_usaha / kode_kbli kosong)[/dim yellow]")
            continue
        tasks.append({
            "index": row_idx,
            "link": str(row[link_col]).strip(),
            "nama_usaha": nama,
            "kode_kbli": kode,
            "komentar": cell(row, komentar_col) or DEFAULT_KOMENTAR,
            "timpa_remark": cell(row, timpa_col),
        })
    return tasks


def load_progress() -> Dict[str, dict]:
    progress: Dict[str, dict] = {}
    if PROGRESS_FILE.exists():
        try:
            with open(PROGRESS_FILE, "r", encoding="utf-8") as f:
                progress = json.load(f)
        except Exception:
            pass

    # Fallback: pulihkan status dari Excel hasil bila JSON hilang
    if RESULT_EXCEL_FILE.exists():
        try:
            wb = openpyxl.load_workbook(RESULT_EXCEL_FILE, data_only=True)
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            header = [str(h).strip().lower() if h is not None else "" for h in (next(rows, None) or [])]

            def col(name):
                return header.index(name) if name in header else None

            c_no, c_link = col("no"), col("link")
            c_kom, c_status, c_detail, c_waktu = col("komentar"), col("status"), col("detail"), col("waktu eksekusi")

            def get(row, c):
                return "" if c is None or len(row) <= c or row[c] is None else str(row[c]).strip()

            for row in rows:
                no = get(row, c_no)
                status = get(row, c_status)
                if not no or not status:
                    continue
                try:
                    no = str(int(float(no)))
                except ValueError:
                    pass
                if no not in progress:
                    progress[no] = {
                        "link": get(row, c_link),
                        "komentar": get(row, c_kom), "status": status,
                        "detail": get(row, c_detail), "timestamp": get(row, c_waktu),
                    }
        except Exception:
            pass
    return progress


def save_progress(data: Dict[str, dict]):
    """Simpan checkpoint secara atomik."""
    PROGRESS_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = PROGRESS_FILE.with_suffix(PROGRESS_FILE.suffix + ".tmp")
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp_path, PROGRESS_FILE)


def export_result_excel(tasks: List[dict], progress_data: Dict[str, dict], output_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hasil Update KBLI"
    ws.append(["No", "Link", "Nama Usaha", "Kode KBLI", "Komentar", "Status", "Detail", "Waktu Eksekusi"])
    for item in tasks:
        rec = progress_data.get(str(item["index"]), {})
        ws.append([
            item["index"], item["link"], item["nama_usaha"], item["kode_kbli"], item["komentar"],
            rec.get("status", "PENDING"), rec.get("detail", ""), rec.get("timestamp", ""),
        ])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


# ─────────────────────────────────────────────────────────────────────────────
# Login SSO
# ─────────────────────────────────────────────────────────────────────────────

def is_page_logged_out(page: Page) -> bool:
    try:
        url = page.url.lower()
        if "sso.bps.go.id" in url or "/login" in url or "/auth/login" in url or "/oauth2/authorization" in url:
            return True
        if page.locator("#kc-form-login, #username, #password, .login-button").count() > 0:
            return True
        text = page.locator("body").inner_text() if page.locator("body").count() > 0 else ""
        return "Selamat Datang Kembali" in text or "Lanjutkan dengan SSO" in text or "Login SSO BPS" in text
    except Exception:
        return False


def perform_sso_login(page: Page, username: str = config.SSO_USERNAME, password: str = config.SSO_PASSWORD) -> bool:
    console.print(f"[bold yellow]🔐 Login otomatis SSO BPS (User: {username})...[/bold yellow]")
    try:
        if "sso.bps.go.id" not in page.url:
            sso_btn = page.locator(
                "a:has-text('Lanjutkan dengan SSO'), a:has-text('Login SSO BPS'), "
                "a[href*='oauth2/authorization'], .login-button"
            )
            if sso_btn.count() > 0 and sso_btn.first.is_visible():
                sso_btn.first.click()
            else:
                page.goto("https://fasih-sm.bps.go.id/oauth2/authorization/ics", wait_until="domcontentloaded")
            page.wait_for_url("**/sso.bps.go.id/**", timeout=20000)

        page.wait_for_selector("#username", timeout=15000)
        page.fill("#username", username)
        page.fill("#password", password)
        page.click("#kc-login")
        page.wait_for_url("**/fasih-sm.bps.go.id/app/**", timeout=25000)
        console.print(f"[bold green]✔ Berhasil login SSO BPS sebagai {username}![/bold green]")
        time.sleep(2.0)
        return True
    except Exception as e:
        console.print(f"[bold red]❌ Gagal login otomatis SSO: {e}[/bold red]")
        return False


# ─────────────────────────────────────────────────────────────────────────────
# Navigasi
# ─────────────────────────────────────────────────────────────────────────────

def _click_menuitem(page: Page, keyword: str) -> bool:
    """Klik item menu sidebar (aside) berdasarkan teks."""
    return bool(page.evaluate(
        r"""(keyword) => {
        const clean = (s) => (s || "").toLowerCase().replace(/\s+/g, ' ').trim();
        const target = clean(keyword);
        const aside = document.querySelector("aside");
        if (!aside) return false;
        const items = Array.from(aside.querySelectorAll("*")).filter(el =>
            el.children.length === 0 && el.textContent.trim().length > 0
        );
        let match = items.find(el => clean(el.textContent) === target);
        if (!match) match = items.find(el => clean(el.textContent).includes(target));
        if (!match) return false;
        let clickable = match;
        while (clickable && clickable !== aside && !clickable.onclick && !String(clickable.className).includes("cursor-pointer")) {
            clickable = clickable.parentElement;
        }
        (clickable || match).scrollIntoView({ block: "center" });
        (clickable || match).click();
        return true;
    }""",
        keyword,
    ))


def _scroll_to_bottom_until(page: Page, script_find: str, step: int = 300, max_steps: int = 40, pause: float = 0.15) -> bool:
    """Gulir bertahap sampai script_find truthy atau mentok. True jika ditemukan."""
    for _ in range(max_steps):
        if page.evaluate(script_find):
            return True
        before = page.evaluate("window.scrollY")
        page.evaluate(_SCROLL_STEP_JS, step)
        time.sleep(pause)
        after = page.evaluate("window.scrollY")
        more = page.evaluate(_HAS_MORE_SCROLL_JS)
        if before == after and not more:
            break
    return bool(page.evaluate(script_find))


def _scroll_to_bottom(page: Page):
    _scroll_to_bottom_until(page, "false")


def _click_usaha_item(page: Page, nama_usaha: str) -> bool:
    """
    Memilih sub-blok usaha pada section KETERANGAN USAHA/PERUSAHAAN.
    Pencocokan teks (bukan sidebar): persis -> diawali -> mengandung -> elemen terpendek yang mengandung.
    """
    target_js = json.dumps(nama_usaha)

    def make_script(do_click: bool) -> str:
        click_line = (
            'best.scrollIntoView({ block: "center" }); best.click(); return true;'
            if do_click else "return true;"
        )
        return f"""(() => {{
        const clean = (s) => (s || "").toLowerCase().replace(/\\s+/g, ' ').trim();
        const target = clean({target_js});
        const visible = (el) => el.offsetWidth > 0 && el.offsetHeight > 0;
        const all = Array.from(document.querySelectorAll("body *")).filter(el => !el.closest("aside"));
        const leaves = all.filter(el => el.children.length === 0 && clean(el.textContent).length > 0);
        const pick = (arr) => arr.find(visible) || arr[0] || null;
        let best = pick(leaves.filter(el => clean(el.textContent) === target))
                || pick(leaves.filter(el => clean(el.textContent).startsWith(target)))
                || pick(leaves.filter(el => clean(el.textContent).includes(target)));
        if (!best) {{
            const any = all.filter(el => clean(el.textContent).includes(target));
            any.sort((x, y) => x.textContent.length - y.textContent.length);
            best = pick(any);
        }}
        if (!best) return false;
        {click_line}
    }})()"""

    _scroll_to_bottom(page)
    if not _scroll_to_bottom_until(page, make_script(do_click=False)):
        return False
    ok = bool(page.evaluate(make_script(do_click=True)))
    if ok:
        time.sleep(1.5)
    return ok


def _select_kbli(page: Page, kode: str) -> Tuple[bool, str]:
    """
    Alur ubah KBLI (sesuai recording):
      DAPATKAN REKOMENDASI KBLI -> Pilih dari Master KBLI -> tombol pertama di wrapper kbli#
      -> textbox 'Pilih salah satu...' diisi kode -> klik opsi '[..][kode]...'.
    Return (ok, wrapper_id_kbli).
    """
    # a. DAPATKAN REKOMENDASI KBLI
    try:
        btn = page.get_by_role("button", name="DAPATKAN REKOMENDASI KBLI").first
        btn.wait_for(state="visible", timeout=10000)
        btn.scroll_into_view_if_needed(timeout=5000)
        btn.click(timeout=8000)
    except Exception:
        return False, "Tombol 'DAPATKAN REKOMENDASI KBLI' tidak ditemukan"
    time.sleep(1.5)

    # b. Pilih dari Master KBLI (rekomendasi bisa butuh waktu memuat)
    try:
        master = page.get_by_text("Pilih dari Master KBLI").first
        master.wait_for(state="visible", timeout=20000)
        master.click(timeout=8000)
    except Exception:
        return False, "Opsi 'Pilih dari Master KBLI' tidak ditemukan"
    time.sleep(1.0)

    # c. Wrapper KBLI -> klik tombol pertama (membuka dropdown)
    wrapper_id = _find_wrapper(page, KBLI_WRAPPER_PREFIX)
    if not wrapper_id:
        return False, f"Field KBLI (prefix {KBLI_WRAPPER_PREFIX}) tidak ditemukan"
    clicked = page.evaluate(
        """(wid) => {
            const w = document.getElementById(wid);
            if (!w) return false;
            const btns = Array.from(w.querySelectorAll("button"));
            const b = btns.find(x => x.offsetWidth > 0 && x.offsetHeight > 0) || btns[0];
            if (!b) return false;
            b.scrollIntoView({ block: "center" });
            b.click();
            return true;
        }""",
        wrapper_id,
    )
    if not clicked:
        return False, f"Tombol dropdown KBLI di wrapper {wrapper_id} tidak ditemukan"

    # d. Textbox 'Pilih salah satu...' (tandai agar bisa dilocate unik)
    mark_js = """(wid) => {
        document.querySelectorAll('[data-auto-kbli]').forEach(e => e.removeAttribute('data-auto-kbli'));
        const all = Array.from(document.querySelectorAll("input, textarea"))
            .filter(e => /pilih salah satu/i.test(e.placeholder || "") && e.offsetWidth > 0);
        const w = document.getElementById(wid);
        const el = all.find(e => w && w.contains(e)) || all[all.length - 1];
        if (!el) return false;
        el.setAttribute('data-auto-kbli', '1');
        return true;
    }"""
    marked = False
    deadline = time.time() + 8
    while time.time() < deadline:
        if page.evaluate(mark_js, wrapper_id):
            marked = True
            break
        time.sleep(0.4)
    if not marked:
        return False, "Textbox 'Pilih salah satu...' tidak muncul"

    try:
        tb = page.locator("[data-auto-kbli='1']").first
        tb.click(timeout=8000)
        tb.fill(kode, timeout=8000)
    except Exception as e:
        return False, f"Gagal mengisi kode KBLI '{kode}': {e}"

    # e. Klik opsi '[..][kode]...' (tunggu hasil pencarian muncul)
    option = None
    deadline = time.time() + 12
    while time.time() < deadline:
        cand = page.get_by_role("option").filter(has_text=f"[{kode}]")
        if cand.count() > 0:
            option = cand.first
            break
        cand = page.get_by_text(f"[{kode}]", exact=False)
        if cand.count() > 0:
            option = cand.first
            break
        time.sleep(0.4)
    if option is None:
        return False, f"Opsi KBLI dengan kode [{kode}] tidak ditemukan"
    try:
        option.scroll_into_view_if_needed(timeout=5000)
        option.click(timeout=8000)
    except Exception as e:
        return False, f"Gagal memilih opsi KBLI [{kode}]: {e}"
    time.sleep(1.0)

    # f. Verifikasi lunak: kode tampil di halaman setelah dipilih
    try:
        page.wait_for_function(
            "(k) => (document.body.innerText || '').includes(k)", arg=kode, timeout=6000
        )
    except Exception:
        console.print(f"   [dim yellow]⚠ Kode {kode} belum terlihat di layar setelah dipilih (verifikasi lunak)[/dim yellow]")

    # wrapper bisa ter-render ulang -> cari lagi id terbarunya
    return True, _find_wrapper(page, KBLI_WRAPPER_PREFIX, timeout=3.0) or wrapper_id


def _find_wrapper(page: Page, wrapper_id: str, timeout: float = 10.0) -> Optional[str]:
    """Cari wrapper dengan id persis (mis. 'no_kk'); fallback ke prefix id. Utamakan yang terlihat."""
    js = """(wid) => {
        const exact = document.getElementById(wid);
        if (exact) return exact.id;
        const els = Array.from(document.querySelectorAll("[id]")).filter(e => e.id.startsWith(wid));
        const chosen = els.find(e => e.offsetWidth > 0 && e.offsetHeight > 0) || els[0] || null;
        return chosen ? chosen.id : null;
    }"""
    deadline = time.time() + timeout
    while time.time() < deadline:
        found = page.evaluate(js, wrapper_id)
        if found:
            return found
        page.evaluate(_SCROLL_STEP_JS, 300)
        time.sleep(0.4)
    return page.evaluate(js, wrapper_id)


# ─────────────────────────────────────────────────────────────────────────────
# Catatan
# ─────────────────────────────────────────────────────────────────────────────

def _close_dialog(page: Page):
    try:
        page.wait_for_selector("[role='dialog']", state="hidden", timeout=3000)
    except Exception:
        page.evaluate(_CLOSE_DIALOG_JS)
        time.sleep(0.4)


def _write_catatan(page: Page, wrapper_id: str, kom_to_fill: str, label: str, timpa: bool) -> Tuple[bool, str]:
    """
    Buka dialog Catatan pada wrapper, baca catatan lama dari textarea, lalu:
      - sama persis dengan yang akan diisi  -> SAME
      - sudah ada #0/#1 dan timpa=False      -> KEPT:<isi lama>
      - selain itu isi/timpa, Simpan, tutup  -> SAVED
    """
    wid = json.dumps(wrapper_id)

    visible_btn_js = f"""(() => {{
        const w = document.getElementById({wid});
        if (!w) return false;
        return Array.from(w.querySelectorAll("button[title='Catatan'], button[aria-label='Catatan']"))
            .some(b => b.offsetWidth > 0 && b.offsetHeight > 0);
    }})()"""
    click_btn_js = f"""(() => {{
        const w = document.getElementById({wid});
        if (!w) return false;
        const btns = Array.from(w.querySelectorAll("button[title='Catatan'], button[aria-label='Catatan']"));
        const b = btns.find(x => x.offsetWidth > 0 && x.offsetHeight > 0) || btns[0];
        if (!b) return false;
        b.scrollIntoView({{ block: "center" }});
        b.click();
        return true;
    }})()"""

    page.evaluate(f"""(() => {{ const w = document.getElementById({wid}); if (w) w.scrollIntoView({{ block: "center" }}); }})()""")
    time.sleep(0.4)
    _scroll_to_bottom_until(page, visible_btn_js, max_steps=10)
    if not page.evaluate(click_btn_js):
        try:   # fallback: getByRole('button', { name: 'Catatan' }) di dalam wrapper
            page.locator(f'[id="{wrapper_id}"]').get_by_role("button", name="Catatan").first.click(timeout=5000)
        except Exception:
            return False, f"Tombol Catatan '{label}' tidak ditemukan"

    try:
        page.wait_for_selector("textarea[placeholder*='catatan' i]", timeout=8000)
    except Exception:
        return False, f"Dialog catatan '{label}' tidak muncul"

    existing = (page.evaluate("""(() => {
        const ta = document.querySelector("textarea[placeholder*='catatan' i]");
        return ta ? (ta.value || "") : "";
    })()""") or "").strip()

    if existing == kom_to_fill:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return True, "SAME"

    if ("#0" in existing or "#1" in existing) and not timpa:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return True, f"KEPT:{existing}"

    res = page.evaluate(
        """(val) => {
            const ta = document.querySelector("textarea[placeholder*='catatan' i]");
            if (!ta) return "ERR";
            ta.focus();
            const setter = Object.getOwnPropertyDescriptor(window.HTMLTextAreaElement.prototype, "value").set;
            setter.call(ta, val);
            ta.dispatchEvent(new Event("input", { bubbles: true }));
            ta.dispatchEvent(new Event("change", { bubbles: true }));
            return ta.value;
        }""",
        kom_to_fill,
    )
    if res != kom_to_fill:
        page.evaluate(_CLOSE_DIALOG_JS)
        _close_dialog(page)
        return False, f"Gagal mengisi textarea catatan '{label}'"

    time.sleep(0.3)
    # Tombol Simpan opsional (di recording catatan langsung ditutup dengan Dismiss)
    simpan_ok = page.evaluate("""(() => {
        const b = document.querySelector("button[title='Simpan']") ||
                  Array.from(document.querySelectorAll("button")).find(x => x.innerText.trim() === "Simpan");
        if (b) { b.click(); return true; }
        return false;
    })()""")
    if not simpan_ok:
        try:   # fallback: getByRole('button', { name: 'Simpan' })
            page.get_by_role("button", name="Simpan").first.click(timeout=3000)
        except Exception:
            pass
    time.sleep(0.6)
    page.evaluate(_CLOSE_DIALOG_JS)   # klik Dismiss
    _close_dialog(page)
    time.sleep(0.4)
    return True, "SAVED"


# ─────────────────────────────────────────────────────────────────────────────
# Submit
# ─────────────────────────────────────────────────────────────────────────────

def submit_assignment(page: Page) -> Tuple[bool, str]:
    """
    Proses submit kuesioner (sama seperti update_penjualan.py / cleaning_nik.py):
    1. Masuk ke section CATATAN
    2. Cek switch "Tampilkan Anomali Usaha dan Keluarga" (aktifkan jika belum)
    3. Pastikan switch "Anomali diselesaikan oleh admin" tidak aktif
    4. Klik Kirim utama -> Cek validasi 0 Galat -> Klik Kirim modal -> Konfirmasi final
    """
    # 1. Masuk ke section CATATAN
    click_catatan_ok = page.evaluate("""(() => {
        const aside = document.querySelector("aside");
        if (!aside) return false;
        const elements = Array.from(aside.querySelectorAll("*")).filter(el => el.children.length === 0 && el.textContent.trim() === "CATATAN");
        if (elements.length === 0) return false;
        const target = elements[0];
        let clickable = target;
        while (clickable && clickable !== aside && !clickable.onclick && !String(clickable.className).includes("cursor-pointer")) {
            clickable = clickable.parentElement;
        }
        (clickable || target).scrollIntoView({ block: "center" });
        (clickable || target).click();
        return true;
    })()""")

    if not click_catatan_ok:
        return False, "Menu CATATAN tidak ditemukan di sidebar"

    time.sleep(1.5)

    # 2. Cek switch di blok CATATAN:
    # - "Tampilkan Anomali Usaha dan Keluarga" jika sudah dichecklist JANGAN dichecklist lagi.
    # - JANGAN checklist "Anomali diselesaikan oleh admin".
    # - JANGAN memberi catatan penjelasan anomali admin.
    page.evaluate("""(() => {
        const rows = Array.from(document.querySelectorAll("div[class*='rounded-lg'], div[class*='gap-2'], div[class*='flex'], label"));

        // Cari switch "Tampilkan Anomali Usaha dan Keluarga" (hindari baris admin)
        const anomaliRow = rows.find(r => {
            const t = (r.innerText || "").toLowerCase();
            return t.includes("tampilkan anomali") && !t.includes("admin");
        });

        if (anomaliRow) {
            const input = anomaliRow.querySelector("input[role='switch'], input[type='checkbox']");
            if (input) {
                const isChecked = input.checked ||
                                  input.getAttribute("aria-checked") === "true" ||
                                  input.hasAttribute("data-checked");
                if (!isChecked) {
                    const control = anomaliRow.querySelector("[id*='control']") || input;
                    control.click();
                }
            }
        }

        // Pastikan switch "Anomali diselesaikan oleh admin" TIDAK DICHECKLIST (jika sempat aktif, uncheck)
        const adminRow = rows.find(r => {
            const t = (r.innerText || "").toLowerCase();
            return t.includes("anomali diselesaikan oleh admin");
        });
        if (adminRow) {
            const adminInput = adminRow.querySelector("input[role='switch'], input[type='checkbox']");
            if (adminInput) {
                const isAdminChecked = adminInput.checked ||
                                       adminInput.getAttribute("aria-checked") === "true" ||
                                       adminInput.hasAttribute("data-checked");
                if (isAdminChecked) {
                    const adminControl = adminRow.querySelector("[id*='control']") || adminInput;
                    adminControl.click();
                }
            }
        }
    })()""")

    time.sleep(1.0)

    # Jeda sebelum klik Kirim (beri waktu catatan & switch tersimpan di server)
    if KIRIM_DELAY > 0:
        console.print(f"   [dim]⏳ Menunggu {KIRIM_DELAY:g} detik sebelum klik Kirim...[/dim]")
        time.sleep(KIRIM_DELAY)

    # 3. Klik tombol Kirim utama (ulangi bila modal validasi belum muncul)
    modal_opened = False
    for _ in range(3):
        page.evaluate("""(() => {
            const btns = Array.from(document.querySelectorAll("button")).filter(b => b.innerText.trim() === "Kirim" && b.offsetWidth > 0);
            const btn = btns[btns.length - 1];
            if (btn) {
                btn.scrollIntoView({ block: "center" });
                btn.click();
            }
        })()""")
        try:
            page.wait_for_selector("[class*='max-w-lg'], [role='dialog']", timeout=6000)
            modal_opened = True
            break
        except Exception:
            time.sleep(1.0)

    if not modal_opened:
        return False, "Modal ringkasan validasi tidak muncul"

    time.sleep(1.0)

    # 4. Validasi 0 Galat & Klik Kirim di modal
    val_res = page.evaluate("""(() => {
        const modal = document.querySelector("[class*='max-w-lg'], [role='dialog']");
        if (!modal) return { status: "ERR", msg: "Modal validasi tidak ditemukan" };
        const modalText = modal.innerText;
        const galatMatch = modalText.match(/GALAT\\s*\\n\\s*(\\d+)/i);
        const galatCount = galatMatch ? parseInt(galatMatch[1]) : 0;
        if (galatCount > 0) {
            const closeBtn = modal.querySelector("button[aria-label='Dismiss']") ||
                             Array.from(modal.querySelectorAll("button")).find(b => b.innerText.trim() === "Batal");
            if (closeBtn) closeBtn.click();
            return { status: "GALAT", count: galatCount, msg: galatCount + " galat ditemukan!" };
        }
        const btns = Array.from(modal.querySelectorAll("button")).filter(b => b.innerText.trim() === "Kirim");
        if (btns.length === 0) return { status: "ERR", msg: "Tombol Kirim di modal validasi tidak ditemukan" };
        btns[0].click();
        return { status: "OK" };
    })()""")

    if val_res.get("status") != "OK":
        return False, f"Validasi gagal: {val_res.get('msg')}"

    # 5. Klik tombol Konfirmasi pada dialog akhir
    try:
        page.wait_for_selector("button:has-text('Konfirmasi')", timeout=10000)
        page.locator("button:has-text('Konfirmasi')").first.click()
    except Exception:
        return False, "Tombol 'Konfirmasi' dialog akhir tidak ditemukan"

    # 6. Verifikasi submit selesai
    try:
        page.wait_for_function("""() => {
            const t = document.body ? document.body.innerText : "";
            return t.includes("Assignment Berhasil Disubmit") || t.includes("Perubahan Anda berhasil disimpan");
        }""", timeout=15000)
        return True, "Assignment Berhasil Disubmit"
    except Exception:
        return False, "Status sukses submit tidak terdeteksi pada layar"


# ─────────────────────────────────────────────────────────────────────────────
# Alur per baris
# ─────────────────────────────────────────────────────────────────────────────

def update_kbli_row(
    page: Page,
    link: str,
    nama_usaha: str,
    kode_kbli: str,
    komentar: str = DEFAULT_KOMENTAR,
    timpa: bool = False,
    submit: bool = True,
    retry_on_logout: bool = True,
) -> Tuple[bool, str]:
    edit_url = link.strip()
    if not edit_url.endswith("/edit"):
        edit_url = edit_url.rstrip("/") + "/edit"

    def _retry():
        return update_kbli_row(page, link, nama_usaha, kode_kbli, komentar, timpa, submit, retry_on_logout=False)

    try:
        page.goto(edit_url, wait_until="domcontentloaded", timeout=30000)
    except Exception as e:
        return False, f"Timeout membuka URL: {e}"
    time.sleep(1.0)

    if is_page_logged_out(page):
        if retry_on_logout and perform_sso_login(page):
            return _retry()
        return False, "Sesi login habis / terlogout"

    try:
        page.wait_for_selector("aside", timeout=20000)
    except Exception:
        if is_page_logged_out(page) and retry_on_logout and perform_sso_login(page):
            return _retry()
        return False, "Menu sidebar (aside) kuesioner tidak muncul"
    time.sleep(1.0)

    # 1. Section KETERANGAN USAHA/PERUSAHAAN
    if not _click_menuitem(page, SECTION_KEYWORD):
        if not _click_menuitem(page, "keterangan usaha"):
            return False, "Section 'KETERANGAN USAHA/PERUSAHAAN' tidak ditemukan di sidebar"
    time.sleep(1.5)

    # 2. Pilih sub-blok usaha
    if not _click_usaha_item(page, nama_usaha):
        return False, f"Sub-blok usaha '{nama_usaha}' tidak ditemukan"
    time.sleep(1.5)

    # 3. Ubah KBLI
    ok, info = _select_kbli(page, kode_kbli)
    if not ok:
        return False, info
    kbli_wrapper = info
    console.print(f"   [dim]↳ KBLI dipilih: [{kode_kbli}] (wrapper {kbli_wrapper})[/dim]")

    # 4. Catatan KBLI (default #1): Catatan -> isi -> Simpan -> Dismiss
    if komentar:
        ok_c, status = _write_catatan(page, kbli_wrapper, komentar, "KBLI", timpa)
        if not ok_c:
            console.print(f"   [dim yellow]⚠ Catatan KBLI gagal: {status}[/dim yellow]")
        elif status.startswith("KEPT"):
            console.print(f"   [dim]↳ Catatan KBLI: catatan lama '{status[5:]}' dipertahankan (timpa_remark bukan 1)[/dim]")
        else:
            console.print(f"   [dim]↳ Catatan KBLI: '{komentar}' [{status}][/dim]")

    # 5. Submit assignment (CATATAN -> switch anomali -> Kirim -> Konfirmasi)
    if not submit:
        return True, f"KBLI [{kode_kbli}] dipilih (tanpa submit, --no-submit)"
    console.print("   [bold cyan]📤 Melakukan Submit Assignment (CATATAN -> Kirim -> Konfirmasi)...[/bold cyan]")
    submit_ok, submit_msg = submit_assignment(page)
    if not submit_ok:
        return False, f"Gagal submit assignment: {submit_msg}"

    console.print(f"   [bold green]✔ Sukses Submit Assignment: {submit_msg}[/bold green]")
    return True, f"KBLI [{kode_kbli}] disimpan, {submit_msg}"


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

def main():
    global KIRIM_DELAY
    parser = argparse.ArgumentParser(description="Batch Automation Update KBLI FASIH-SM SE2026 (1-Command)")
    parser.add_argument("--source", default=str(DEFAULT_SOURCE_FILE), help="File Excel sumber (default: source/update_kbli.xlsx)")
    parser.add_argument("--init-sample", action="store_true", help="Buat file Excel sumber contoh lalu keluar")
    parser.add_argument("--start", type=int, default=1, help="Nomor baris awal (default: 1)")
    parser.add_argument("--end", type=int, default=None, help="Nomor baris akhir")
    parser.add_argument("--limit", type=int, default=None, help="Jumlah baris maksimum yang diproses")
    parser.add_argument("--delay", type=float, default=1.5, help="Jeda antar baris dalam detik (default: 1.5)")
    parser.add_argument("--delay-kirim", type=float, default=KIRIM_DELAY, help="Jeda sebelum klik Kirim dalam detik (default: 10)")
    parser.add_argument("--force", action="store_true", help="Proses ulang semua baris meskipun sudah SUCCESS")
    parser.add_argument("--headless", action="store_true", help="Jalankan tanpa jendela Chrome")
    parser.add_argument("--dry-run", action="store_true", help="Hanya tampilkan daftar baris tanpa eksekusi")
    parser.add_argument("--no-submit", action="store_true", help="Uji alur tanpa submit (perubahan bisa tidak tersimpan)")
    args = parser.parse_args()

    KIRIM_DELAY = args.delay_kirim

    source_path = Path(args.source)

    if args.init_sample:
        create_sample_source(source_path)
        return

    all_tasks = read_excel_source(source_path)
    progress_data = load_progress()

    completed = {int(k) for k, v in progress_data.items() if v.get("status") == "SUCCESS"}
    failed = {int(k) for k, v in progress_data.items() if v.get("status") == "FAILED"}

    filtered_tasks = []
    skipped = 0
    for t in all_tasks:
        idx = t["index"]
        if idx < args.start:
            continue
        if args.end and idx > args.end:
            continue
        if not args.force and idx in completed:
            skipped += 1
            continue
        filtered_tasks.append(t)
        if args.limit and len(filtered_tasks) >= args.limit:
            break

    console.print(
        Panel.fit(
            f"[bold cyan]⚡ FASIH-SM Batch Update KBLI Engine[/bold cyan]\n"
            f"File Sumber      : [yellow]{source_path}[/yellow] (Total {len(all_tasks)} baris)\n"
            f"SSO User         : [bold green]{config.SSO_USERNAME}[/bold green] (Auto-Login: {config.AUTO_LOGIN_SSO})\n"
            f"Profil Browser   : [magenta]{config.BROWSER_DATA_DIR}[/magenta]\n"
            f"Sudah Sukses     : [bold green]{len(completed)} baris[/bold green] (Otomatis dilewati: {skipped})\n"
            f"Pernah Gagal     : [bold red]{len(failed)} baris[/bold red] (Akan dicoba ulang)\n"
            f"Target Proses Ini: [bold cyan]{len(filtered_tasks)} baris[/bold cyan] (Mulai baris {args.start})\n"
            f"Tampilan Visual  : [yellow]{'Headless' if args.headless else 'Google Chrome Jendela Visual'}[/yellow] | Delay: [cyan]{args.delay}s[/cyan]",
            border_style="cyan",
        )
    )

    if not filtered_tasks:
        console.print(
            "[bold green]✔ Semua baris target sudah SUCCESS sebelumnya![/bold green]\n"
            "[dim]Gunakan --force jika ingin memproses ulang.[/dim]"
        )
        return

    if args.dry_run:
        table = Table(title="Daftar Baris Target (Dry Run)")
        table.add_column("No", style="cyan", justify="right")
        table.add_column("Nama Usaha", style="magenta", overflow="fold")
        table.add_column("Kode KBLI", style="green")
        table.add_column("Komentar", style="yellow")
        table.add_column("Timpa", style="red")
        table.add_column("Status Terakhir", style="yellow")
        for t in filtered_tasks[:20]:
            prev = progress_data.get(str(t["index"]), {}).get("status", "PENDING")
            table.add_row(str(t["index"]), t["nama_usaha"], t["kode_kbli"], t["komentar"] or "-",
                          "ya" if _timpa_kbli(t["timpa_remark"]) else "tidak", prev)
        console.print(table)
        if len(filtered_tasks) > 20:
            console.print(f"[dim]... dan {len(filtered_tasks) - 20} baris lainnya.[/dim]")
        return

    console.print(f"[cyan]🚀 Meluncurkan Google Chrome otomatis ({'headless' if args.headless else 'visual'})...[/cyan]")
    config.BROWSER_DATA_DIR.mkdir(parents=True, exist_ok=True)

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(config.BROWSER_DATA_DIR),
            channel=config.BROWSER_CHANNEL,
            headless=args.headless,
            viewport=None,
            args=["--start-maximized", "--disable-blink-features=AutomationControlled"],
        )
        page = context.pages[0] if context.pages else context.new_page()

        try:
            page.goto(f"{config.FASIH_BASE_URL}/app/", wait_until="domcontentloaded", timeout=25000)
            time.sleep(2.0)
            if is_page_logged_out(page):
                console.print("[yellow]⚠ Status awal logout! Login otomatis SSO BPS...[/yellow]")
                if not perform_sso_login(page):
                    console.print("[bold red]❌ Gagal login otomatis SSO. Proses dibatalkan.[/bold red]")
                    context.close()
                    return
        except Exception as e:
            console.print(f"[dim yellow]Peringatan cek awal: {e}[/dim yellow]")

        success_count = 0
        fail_count = 0
        current_item = None

        def record(item: dict, status: str, detail: str):
            progress_data[str(item["index"])] = {
                "link": item["link"], "nama_usaha": item["nama_usaha"], "kode_kbli": item["kode_kbli"],
                "komentar": item["komentar"], "status": status, "detail": detail,
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            }
            save_progress(progress_data)

        try:
            for i, item in enumerate(filtered_tasks, start=1):
                current_item = item
                idx = item["index"]

                console.print(f"\n[bold blue]━━━ [{i}/{len(filtered_tasks)}] Memproses Baris #{idx} ━━━[/bold blue]")
                console.print(f"Usaha   : [magenta]{item['nama_usaha']}[/magenta] | KBLI: [cyan]{item['kode_kbli']}[/cyan] | Komentar: [cyan]{item['komentar'] or '-'}[/cyan]")
                console.print(f"Link    : [dim]{item['link']}[/dim]")

                try:
                    is_ok, msg = update_kbli_row(
                        page, item["link"], item["nama_usaha"], item["kode_kbli"], item["komentar"],
                        timpa=_timpa_kbli(item["timpa_remark"]), submit=not args.no_submit,
                    )
                except KeyboardInterrupt:
                    raise
                except Exception as e:
                    is_ok, msg = False, f"Exception tak terduga: {e}"

                record(item, "SUCCESS" if is_ok else "FAILED", msg)
                current_item = None

                if is_ok:
                    success_count += 1
                    console.print(f"[bold green]✔ #{idx} Berhasil: {msg}[/bold green]")
                else:
                    fail_count += 1
                    console.print(f"[bold red]❌ #{idx} Gagal: {msg}[/bold red]")

                if i % 10 == 0 or i == len(filtered_tasks):
                    export_result_excel(all_tasks, progress_data, RESULT_EXCEL_FILE)

                if i < len(filtered_tasks) and args.delay > 0:
                    time.sleep(args.delay)

        except KeyboardInterrupt:
            console.print("\n[bold yellow]⚠ Eksekusi dihentikan pengguna (Ctrl+C). Checkpoint tersimpan aman![/bold yellow]")
            if current_item is not None:
                record(current_item, "FAILED", "Dihentikan pengguna (Ctrl+C) saat sedang diproses")
                console.print(f"[yellow]↳ Baris #{current_item['index']} ditandai FAILED (terhenti) & akan dicoba ulang.[/yellow]")
        finally:
            export_result_excel(all_tasks, progress_data, RESULT_EXCEL_FILE)
            context.close()
            console.print(
                Panel.fit(
                    f"[bold cyan]📊 Rekapitulasi Update KBLI[/bold cyan]\n"
                    f"Selesai Diproses : [bold]{success_count + fail_count}[/bold] baris\n"
                    f"Berhasil (Submit): [bold green]{success_count}[/bold green]\n"
                    f"Gagal            : [bold red]{fail_count}[/bold red]\n"
                    f"Log Checkpoint   : [yellow]{PROGRESS_FILE}[/yellow]\n"
                    f"Hasil Excel      : [green]{RESULT_EXCEL_FILE}[/green]",
                    border_style="cyan",
                )
            )


if __name__ == "__main__":
    main()